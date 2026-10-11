"""Parse uploaded CSV / XLSX content into rows of text, keyed by normalised header."""

import csv
import io
import re
import zipfile
from typing import Any

from openpyxl import load_workbook

from app.core.exceptions import UnprocessableError, UnsupportedMediaTypeError

MAX_IMPORT_ROWS = 5_000
MAX_COLUMNS = 50
#: Guards against zip bombs: total uncompressed size of an uploaded workbook.
MAX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024


def normalise_header(header: Any) -> str:
    text = re.sub(r"[^0-9a-zA-Z]+", "_", str(header or "").strip().lower())
    return text.strip("_")


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if hasattr(value, "date") and hasattr(value, "hour"):  # datetime
        return value.date().isoformat()
    return str(value).strip()


def parse_table(extension: str, content: bytes) -> list[dict[str, str]]:
    if extension == ".csv":
        rows = _read_csv(content)
    elif extension == ".xlsx":
        rows = _read_xlsx(content)
    else:
        raise UnsupportedMediaTypeError("only .csv and .xlsx files can be imported")
    if len(rows) < 2:
        raise UnprocessableError(
            "the file needs a header row and at least one data row"
        )
    headers = [normalise_header(h) for h in rows[0]]
    if len(headers) > MAX_COLUMNS:
        raise UnprocessableError(f"too many columns (maximum {MAX_COLUMNS})")
    named = [h for h in headers if h]
    if len(set(named)) != len(named):
        raise UnprocessableError("the header row has duplicate column names")
    data = rows[1:]
    if len(data) > MAX_IMPORT_ROWS:
        raise UnprocessableError(f"too many rows (maximum {MAX_IMPORT_ROWS})")
    result = []
    for row in data:
        item = {h: _cell_text(v) for h, v in zip(headers, row, strict=False) if h}
        if any(item.values()):  # skip fully blank lines
            result.append(item)
    return result


def _read_csv(content: bytes) -> list[list[str]]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise UnsupportedMediaTypeError("CSV files must be UTF-8 encoded") from exc
    header_line = text.lstrip("\r\n").split("\n", 1)[0]
    delimiter = max(",;\t", key=header_line.count)
    reader = csv.reader(io.StringIO(text, newline=""), delimiter=delimiter)
    rows: list[list[str]] = []
    for row in reader:
        rows.append(row)
        if len(rows) > MAX_IMPORT_ROWS + 1:
            raise UnprocessableError(f"too many rows (maximum {MAX_IMPORT_ROWS})")
    return rows


def _read_xlsx(content: bytes) -> list[list[Any]]:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if (
                sum(info.file_size for info in archive.infolist())
                > MAX_UNCOMPRESSED_BYTES
            ):
                raise UnprocessableError("the workbook is too large once unpacked")
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except (zipfile.BadZipFile, KeyError, ValueError, OSError) as exc:
        raise UnsupportedMediaTypeError(
            "the file is not a valid .xlsx workbook"
        ) from exc
    try:
        sheet = workbook.worksheets[0]
        rows: list[list[Any]] = []
        for row in sheet.iter_rows(values_only=True):
            rows.append(list(row))
            if len(rows) > MAX_IMPORT_ROWS + 1:
                raise UnprocessableError(f"too many rows (maximum {MAX_IMPORT_ROWS})")
        return rows
    finally:
        workbook.close()
