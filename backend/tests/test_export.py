import io
import re
from datetime import UTC, datetime, timedelta
from decimal import Decimal
from types import SimpleNamespace
from typing import Any

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pydantic import ValidationError

from app.api import providers
from app.api.dependencies import get_authenticated_user
from app.core.authz import AuthContext
from app.core.exceptions import UnprocessableError
from app.exports.base import ExportContext, Scope
from app.exports.formatting import format_value
from app.exports.model import ExportColumn
from app.exports.pdf import render_pdf
from app.exports.tickets import TicketsExport, prepare
from app.exports.xlsx import render_xlsx
from app.main import app
from app.models.auth import User
from app.models.category import RequestCategory
from app.models.equipment import Equipment, EquipmentType
from app.models.maintenance import (
    MaintenanceRecord,
    MaintenanceRequest,
    TicketAssignment,
)
from app.models.priority import Priority
from app.models.technician import Technician
from app.repositories.export import ExportRepository
from app.schemas.export import ExportFilters, ExportRequest
from app.services import export as export_module
from app.services.export import ExportService, filename

NOW = datetime(2026, 10, 14, 15, 0, tzinfo=UTC)
ADMIN = AuthContext(
    user=SimpleNamespace(id=1, full_name="Ana Pérez"),
    roles=frozenset({"ADMIN"}),
    permissions=frozenset({"reports.read"}),
)
TECHNICIAN = AuthContext(
    user=SimpleNamespace(id=7, full_name="Luis Gómez"),
    roles=frozenset({"TECHNICIAN"}),
    permissions=frozenset({"reports.read"}),
)


def aware(day: int, month: int = 9, year: int = 2026) -> str:
    return f"{year}-{month:02d}-{day:02d}T00:00:00Z"


def raw_row(n: int, **over: Any) -> dict[str, Any]:
    """A row shaped like the repository returns it."""
    created = NOW - timedelta(days=n, hours=3)
    done = n % 3 == 0
    row = {
        "id": n,
        "ticket_number": f"TCK-2026-{n:06d}",
        "title": f"Pantalla rota {n} con una descripción bastante larga para ajustar",
        "status": "resolved" if done else "in_progress",
        "priority_code": "high" if n % 2 else "low",
        "priority": "Alta" if n % 2 else "Baja",
        "sla_hours": 24 if n % 2 else None,
        "category_code": "hardware",
        "category": "Hardware",
        "equipment": f"LT-{n:03d}",
        "serial": f"SN{n}",
        "equipment_type": "Laptop",
        "requested_by": "María",
        "technician_id": 1 if n % 4 else None,
        "technician": "Luis Gómez" if n % 4 else None,
        "created_at": created,
        "resolved_at": created + timedelta(hours=10 + n) if done else None,
        "labor_cost": Decimal("100.50") if done else Decimal(0),
        "parts_cost": Decimal("49.50") if done else Decimal(0),
    }
    row.update(over)
    return row


# -- formatting ------------------------------------------------------------------------------


def test_values_are_shown_the_spanish_way():
    assert format_value(Decimal("1234.5"), "money") == "$ 1.234,50"
    assert format_value(12.34, "number") == "12,3"
    assert format_value(1234567, "int") == "1.234.567"
    assert format_value(NOW, "date") == "14/10/2026"  # 10:00 in UTC-5
    assert (
        format_value(datetime(2026, 10, 15, 2, 0, tzinfo=UTC), "date") == "14/10/2026"
    )
    assert format_value(NOW, "datetime") == "14/10/2026 10:00"
    assert format_value(True, "bool") == "Sí"
    assert format_value(None, "bool") == "N/A"
    assert format_value(None, "date") == "—"


# -- the ticket report ----------------------------------------------------------------------


def test_prepare_derives_hours_sla_and_cost():
    done = prepare(raw_row(3))
    assert done["resolution_hours"] == 13.0
    assert done["sla_met"] is True  # high priority, 24 h target
    assert done["total_cost"] == Decimal("150.00")
    assert done["status"] == "Resuelto"
    late = prepare(raw_row(30, status="resolved", resolved_at=NOW, sla_hours=24))
    assert late["sla_met"] is False
    open_ticket = prepare(raw_row(1))
    assert open_ticket["resolution_hours"] is None and open_ticket["sla_met"] is None
    no_sla = prepare(raw_row(6))  # low priority has no target
    assert no_sla["sla_met"] is None
    uncategorised = prepare(raw_row(1, category=None, category_code=None))
    assert uncategorised["category"] == "Sin categoría"


def context(**filters: Any) -> ExportContext:
    return ExportContext(
        company="Mantix S.A.S.",
        legal_name="Mantix Servicios S.A.S.",
        tax_id="900.123.456-7",
        generated_by="Ana Pérez",
        generated_at=NOW,
        filters=ExportFilters(**filters),
    )


def test_the_report_describes_itself_and_summarises_the_rows():
    report = TicketsExport().build(
        [raw_row(n) for n in range(1, 13)],
        context(statuses=["resolved"], date_from=aware(1), date_to=aware(30)),
    )
    left = {m.label: m.value for m in report.meta_left}
    right = {m.label: m.value for m in report.meta_right}
    assert left["Empresa"] == "Mantix S.A.S." and left["NIT"] == "900.123.456-7"
    assert left["Generado por"] == "Ana Pérez"
    assert right["Período"] == "31/08/2026 – 29/09/2026" or right["Período"].startswith(
        "0"
    )
    assert right["Total de registros"] == "12"
    assert right["Filtros"] == "Estado: Resuelto"
    groups = {
        g.title: {line.label: line.value for line in g.lines} for g in report.summary
    }
    assert groups["Por estado"] == {"En curso": 8, "Resuelto": 4}
    assert groups["Por prioridad"] == {"Alta": 6, "Baja": 6}
    assert groups["Totales"]["Total de tickets"] == 12
    assert groups["Totales"]["Costo total"] == pytest.approx(600.0)
    technicians = {r["technician"]: r for r in report.extra_sheets[0].rows}
    assert technicians["Luis Gómez"]["tickets"] == 9
    assert technicians["Sin asignar"]["tickets"] == 3
    assert list(technicians)[-1] == "Sin asignar"


def test_totals_are_none_when_nothing_is_measurable():
    report = TicketsExport().build([raw_row(1)], context())
    totals = {line.label: line for line in report.summary[-1].lines}
    assert totals["Tiempo medio de resolución (h)"].value is None
    assert totals["Tiempo medio de resolución (h)"].text == "N/A"
    assert totals["Cumplimiento de SLA (%)"].text == "N/A"


# -- Excel -----------------------------------------------------------------------------------


def workbook(rows=24, **options):
    report = TicketsExport().build([raw_row(n) for n in range(1, rows + 1)], context())
    return report, load_workbook(io.BytesIO(render_xlsx(report, **options)))


def test_excel_has_summary_detail_and_technician_sheets():
    _, book = workbook()
    assert book.sheetnames == ["Resumen", "Detalle", "Por técnico"]
    summary = book["Resumen"]
    assert len(summary._images) == 1 and len(summary._charts) == 2
    assert summary["C1"].value == "REPORTE DE TICKETS"
    labels = [c.value for row in summary.iter_rows(min_row=1, max_row=40) for c in row]
    assert {"Por estado", "Por prioridad", "Totales"} <= set(labels)


def test_excel_detail_is_a_real_table_with_formats_filters_and_formulas():
    report, book = workbook(rows=24)
    sheet = book["Detalle"]
    header = 8
    assert [c.value for c in sheet[header]][:3] == [
        "Ticket",
        "Título",
        "Código de equipo",
    ]
    assert sheet.freeze_panes == "A9"
    assert sheet.auto_filter.ref == f"A8:{chr(64 + len(report.columns))}32"
    assert len(sheet._images) == 1
    headers = [c.value for c in sheet[header]]
    created = sheet.cell(row=9, column=headers.index("Creado") + 1)
    assert isinstance(created.value, datetime) and created.number_format == "dd/mm/yyyy"
    cost = sheet.cell(row=9, column=headers.index("Costo total") + 1)
    assert isinstance(cost.value, (int, float)) and "$" in cost.number_format
    hours = sheet.cell(row=11, column=headers.index("Tiempo de resolución (h)") + 1)
    assert isinstance(hours.value, (int, float))
    total = headers.index("Costo total") + 1
    letter = chr(64 + total)
    assert (
        sheet.cell(row=33, column=total).value == f"=SUBTOTAL(109,{letter}9:{letter}32)"
    )
    assert sheet["A33"].value == '="Total: "&SUBTOTAL(103,A9:A32)'
    assert sheet.cell(row=9, column=headers.index("SLA cumplido") + 1).value in {
        "Sí",
        "No",
        None,
    }


def test_excel_columns_can_be_chosen_and_charts_left_out():
    report = TicketsExport().build([raw_row(n) for n in range(1, 6)], context())
    chosen = report.only(["ticket_number", "total_cost"])
    book = load_workbook(io.BytesIO(render_xlsx(chosen, include_charts=False)))
    assert [c.value for c in book["Detalle"][8] if c.value] == ["Ticket", "Costo total"]
    assert len(book["Resumen"]._charts) == 0
    without = load_workbook(io.BytesIO(render_xlsx(chosen, include_summary=False)))
    assert without.sheetnames == ["Detalle", "Por técnico"]


def test_excel_never_evaluates_user_text_as_a_formula():
    evil = raw_row(1, title='=HYPERLINK("http://x")', technician="+cmd|' /C calc'!A0")
    report = TicketsExport().build([evil], context())
    sheet = load_workbook(io.BytesIO(render_xlsx(report)))["Detalle"]
    titles = [c.value for c in sheet[9]]
    assert '\'=HYPERLINK("http://x")' in titles
    assert not any(isinstance(v, str) and v.startswith("=") for v in titles)


# -- PDF -------------------------------------------------------------------------------------


def pdf_pages(content: bytes) -> int:
    assert content.startswith(b"%PDF-") and content.rstrip().endswith(b"%%EOF")
    return int(re.findall(rb"/Count (\d+)", content)[-1])


def test_pdf_opens_and_paginates_long_tables():
    report = TicketsExport().build([raw_row(n) for n in range(1, 121)], context())
    short = TicketsExport().build([raw_row(n) for n in range(1, 4)], context())
    assert pdf_pages(render_pdf(report.only(list(TicketsExport.pdf_default)))) > 3
    assert (
        pdf_pages(
            render_pdf(
                short.only(list(TicketsExport.pdf_default)), include_summary=False
            )
        )
        == 1
    )


def test_pdf_summary_and_charts_are_optional():
    report = TicketsExport().build([raw_row(n) for n in range(1, 4)], context())
    chosen = report.only(["ticket_number", "status"])
    full = render_pdf(chosen)
    no_charts = render_pdf(chosen, include_charts=False)
    no_summary = render_pdf(chosen, include_summary=False)
    assert len(full) > len(no_charts) > len(no_summary)
    assert pdf_pages(no_summary) == 1


def test_pdf_handles_an_empty_table_and_odd_text():
    odd = raw_row(1, title='<b>Ñandú & <script> "quote" áéíóú</b>')
    report = TicketsExport().build([odd], context())
    assert pdf_pages(render_pdf(report)) >= 1
    empty = TicketsExport().build([], context())
    assert pdf_pages(render_pdf(empty)) >= 1
    assert isinstance(ExportColumn("a", "A").numeric, bool)


# -- request validation --------------------------------------------------------------------


@pytest.mark.parametrize(
    "filters",
    [
        {"date_from": aware(10), "date_to": aware(9)},  # inverted
        {"date_from": aware(1, 1, 2025), "date_to": aware(1, 3, 2026)},  # over 366 days
        {"date_from": "2026-09-01T00:00:00"},  # naive
        {"date_from": "2999-01-01T00:00:00Z"},  # future
        {"statuses": ["Bad Value"]},
        {"technician_ids": [0]},
        {"technician_ids": ["x"]},
        {"unknown": ["x"]},
        {"statuses": [f"s{i}x" for i in range(21)]},
    ],
)
def test_filters_reject_incoherent_values(filters):
    with pytest.raises(ValidationError):
        ExportFilters(**filters)


def test_filters_accept_the_boundary_and_drop_repeats():
    exact = ExportFilters(date_from=aware(1, 1, 2026), date_to=aware(2, 1, 2027))
    assert exact.date_to - exact.date_from == timedelta(days=366)
    assert (
        ExportFilters(statuses=["ab", "ab"], technician_ids=[]).technician_ids is None
    )


def test_request_needs_at_least_one_column_and_no_extras():
    with pytest.raises(ValidationError):
        ExportRequest(columns=[])
    with pytest.raises(ValidationError):
        ExportRequest(columns=["Bad Key"])
    with pytest.raises(ValidationError):
        ExportRequest(nonsense=1)
    assert ExportRequest(columns=["a_b", "a_b"]).columns == ["a_b"]


# -- service (fake repository) ---------------------------------------------------------------


class FakeAudit:
    def __init__(self):
        self.records: list[dict[str, Any]] = []

    def record(self, **kwargs):
        self.records.append(kwargs)


class FakeRepository:
    def __init__(self, rows=None, *, organization=None, technician_id=None):
        self.rows = rows if rows is not None else [raw_row(n) for n in range(1, 9)]
        self._organization = organization
        self._technician_id = technician_id
        self.scopes: list[tuple[Any, Any]] = []
        self.runs: list[dict[str, Any]] = []
        self.committed = False

    async def organization(self):
        return self._organization

    async def technician_id_of(self, user_id):
        return self._technician_id

    async def count_tickets(self, filters, *, requester_id=None, technician_id=None):
        self.scopes.append((requester_id, technician_id))
        return len(self.rows)

    async def export_tickets(
        self, filters, *, requester_id=None, technician_id=None, limit
    ):
        return self.rows[:limit]

    def record_run(self, **kwargs):
        self.runs.append(kwargs)
        return SimpleNamespace(id=len(self.runs))

    async def flush(self):
        pass

    async def commit(self):
        self.committed = True


def make_service(**kwargs):
    repo = FakeRepository(**kwargs)
    return ExportService(repo, FakeAudit()), repo


async def test_exports_a_pdf_and_an_excel_with_a_dated_name_and_audit():
    service, repo = make_service()
    pdf = await service.export(ADMIN, "tickets", "pdf", ExportRequest(), now=NOW)
    assert pdf.filename == "mantix_tickets_2026-10-14.pdf"
    assert pdf.media_type == "application/pdf" and pdf_pages(pdf.content) >= 1
    xlsx = await service.export(ADMIN, "tickets", "xlsx", ExportRequest(), now=NOW)
    assert xlsx.filename == "mantix_tickets_2026-10-14.xlsx"
    assert load_workbook(io.BytesIO(xlsx.content)).sheetnames[0] == "Resumen"
    audit = service.audit.records[0]
    assert audit["action"] == "report.exported"
    assert audit["new_values"]["rows"] == 8 and audit["new_values"]["format"] == "pdf"
    assert repo.runs[0]["row_count"] == 8 and repo.committed


async def test_the_file_name_uses_display_time_not_utc():
    late = datetime(2026, 10, 15, 2, 0, tzinfo=UTC)  # still the 14th in UTC-5
    assert filename("tickets", "pdf", late) == "mantix_tickets_2026-10-14.pdf"


async def test_defaults_are_a_compact_pdf_and_a_complete_excel():
    service, _ = make_service()
    pdf_keys = service._columns(TicketsExport(), "pdf", None)
    xlsx_keys = service._columns(TicketsExport(), "xlsx", None)
    assert pdf_keys == list(TicketsExport.pdf_default)
    assert len(xlsx_keys) == len(TicketsExport.columns) > len(pdf_keys)


async def test_columns_are_validated_and_kept_in_report_order():
    service, _ = make_service()
    with pytest.raises(UnprocessableError, match="unknown columns: nope"):
        await service.export(
            ADMIN, "tickets", "xlsx", ExportRequest(columns=["nope"]), now=NOW
        )
    out = await service.export(
        ADMIN,
        "tickets",
        "xlsx",
        ExportRequest(columns=["total_cost", "ticket_number"]),
        now=NOW,
    )
    headers = [
        c.value for c in load_workbook(io.BytesIO(out.content))["Detalle"][8] if c.value
    ]
    assert headers == ["Ticket", "Costo total"]


async def test_nothing_to_export_and_too_much_are_422_with_a_reason():
    service, _ = make_service(rows=[])
    with pytest.raises(UnprocessableError, match="nothing to export"):
        await service.export(ADMIN, "tickets", "pdf", ExportRequest(), now=NOW)
    service, repo = make_service()

    async def many(*args, **kwargs):
        return export_module.MAX_EXPORT_ROWS + 1

    repo.count_tickets = many
    with pytest.raises(UnprocessableError, match="narrow the filters"):
        await service.export(ADMIN, "tickets", "pdf", ExportRequest(), now=NOW)


async def test_reports_without_an_export_yet_say_so():
    service, _ = make_service()
    with pytest.raises(UnprocessableError, match="cannot be exported yet"):
        await service.export(ADMIN, "equipment", "pdf", ExportRequest(), now=NOW)
    with pytest.raises(UnprocessableError):
        service.options("inventory_items")


async def test_each_role_exports_only_what_it_may_see():
    admin_service, admin_repo = make_service()
    await admin_service.count(ADMIN, "tickets", ExportFilters())
    assert admin_repo.scopes == [(None, None)]
    tech_service, tech_repo = make_service(technician_id=5)
    await tech_service.count(TECHNICIAN, "tickets", ExportFilters())
    assert tech_repo.scopes == [(None, 5)]
    other_service, other_repo = make_service(technician_id=None)
    await other_service.count(TECHNICIAN, "tickets", ExportFilters())
    assert other_repo.scopes == [(7, None)]
    assert isinstance(Scope(), Scope)


async def test_the_header_uses_the_company_settings_or_falls_back():
    service, _ = make_service()
    default = await service.export(ADMIN, "tickets", "xlsx", ExportRequest(), now=NOW)
    sheet = load_workbook(io.BytesIO(default.content))["Detalle"]
    values = {c.value for row in sheet.iter_rows(min_row=4, max_row=7) for c in row}
    assert "Mantix" in values and "900.123.456-7" not in values
    org = SimpleNamespace(
        company_name="Majesa", legal_name="Majesa S.A.S.", tax_id="900.111.222-3"
    )
    service, _ = make_service(organization=org)
    custom = await service.export(ADMIN, "tickets", "xlsx", ExportRequest(), now=NOW)
    sheet = load_workbook(io.BytesIO(custom.content))["Detalle"]
    values = {c.value for row in sheet.iter_rows(min_row=4, max_row=7) for c in row}
    assert {"Majesa", "Majesa S.A.S.", "900.111.222-3"} <= values


async def test_the_options_list_the_real_columns_with_their_defaults():
    service, _ = make_service()
    options = service.options("tickets")
    keys = {c.key: c for c in options.columns}
    assert keys["serial"].default_xlsx and not keys["serial"].default_pdf
    assert keys["total_cost"].default_pdf
    assert options.formats == ["pdf", "xlsx"] and options.max_rows == 10_000


# -- routes -----------------------------------------------------------------------------------


def as_user(permissions, *, service=None):
    user = SimpleNamespace(
        id=1,
        full_name="Ana Pérez",
        roles=[
            SimpleNamespace(
                code="ADMIN",
                is_active=True,
                permissions=[
                    SimpleNamespace(code=c, is_active=True) for c in permissions
                ],
            )
        ],
    )
    app.dependency_overrides[get_authenticated_user] = lambda: user
    if service is not None:
        app.dependency_overrides[providers.get_export_service] = lambda: service


@pytest.fixture
def client():
    service, _ = make_service()
    as_user(["reports.read"], service=service)
    with TestClient(app, raise_server_exceptions=False) as test_client:
        yield test_client
    app.dependency_overrides.clear()


def test_the_export_endpoint_streams_the_file_with_its_name(client):
    response = client.post("/reports/tickets/export?format=pdf", json={})
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/pdf"
    assert re.match(
        r'attachment; filename="mantix_tickets_\d{4}-\d{2}-\d{2}\.pdf"',
        response.headers["content-disposition"],
    )
    xlsx = client.post(
        "/reports/tickets/export?format=xlsx", json={"include_charts": False}
    )
    assert xlsx.status_code == 200 and load_workbook(io.BytesIO(xlsx.content))


def test_invalid_requests_are_422_with_a_specific_message(client):
    inverted = client.post(
        "/reports/tickets/export?format=pdf",
        json={"filters": {"date_from": aware(10), "date_to": aware(1)}},
    )
    assert inverted.status_code == 422
    assert "date_to cannot be before date_from" in inverted.text
    assert client.post("/reports/tickets/export?format=csv", json={}).status_code == 422
    assert (
        client.post(
            "/reports/tickets/export?format=pdf", json={"columns": []}
        ).status_code
        == 422
    )
    assert (
        client.post("/reports/nonsense/export?format=pdf", json={}).status_code == 422
    )
    unsupported = client.post("/reports/equipment/export?format=pdf", json={})
    assert (
        unsupported.status_code == 422 and "cannot be exported yet" in unsupported.text
    )


def test_count_and_options_endpoints(client):
    count = client.post(
        "/reports/tickets/export/count", json={"statuses": ["resolved"]}
    ).json()
    assert count == {
        "total": 8,
        "max_rows": 10000,
        "warn_rows": 2000,
        "over_limit": False,
        "large": False,
    }
    options = client.get("/reports/tickets/export/options").json()
    assert (
        options["report"] == "tickets"
        and options["columns"][0]["key"] == "ticket_number"
    )


def test_exporting_needs_the_reports_permission():
    service, _ = make_service()
    as_user([], service=service)
    try:
        with TestClient(app, raise_server_exceptions=False) as test_client:
            for call in (
                lambda: test_client.post("/reports/tickets/export?format=pdf", json={}),
                lambda: test_client.post("/reports/tickets/export/count", json={}),
                lambda: test_client.get("/reports/tickets/export/options"),
            ):
                assert call().status_code == 403
    finally:
        app.dependency_overrides.clear()


# -- real SQL (needs TEST_DATABASE_URL) -----------------------------------------------------------


async def seed_tickets(session):
    kind = EquipmentType(code="LAP", name="Laptop", is_active=True)
    boss = User(clerk_user_id="c1", email="a@x.co", full_name="Ana", is_active=True)
    tech_user = User(
        clerk_user_id="c2", email="l@x.co", full_name="Luis", is_active=True
    )
    other_user = User(
        clerk_user_id="c3", email="o@x.co", full_name="Olga", is_active=True
    )
    session.add_all([kind, boss, tech_user, other_user])
    await session.flush()
    technician = Technician(user_id=tech_user.id, is_available=True)
    equipment = Equipment(
        asset_tag="LT-1",
        serial_number="SN-1",
        equipment_type_id=kind.id,
        brand="Dell",
        model="X",
        status="assigned",
    )
    session.add_all(
        [
            technician,
            equipment,
            Priority(code="high", name="Alta", level=3, sla_hours=24, is_active=True),
            RequestCategory(code="hardware", name="Hardware", is_active=True),
        ]
    )
    await session.flush()

    def ticket(
        n, requester, status="resolved", created_days=5, hours=10, category="hardware"
    ):
        created = NOW - timedelta(days=created_days)
        return MaintenanceRequest(
            ticket_number=f"T-{n}",
            equipment_id=equipment.id,
            requested_by_user_id=requester.id,
            title=f"t{n}",
            description="d",
            priority="high",
            category=category,
            status=status,
            created_at=created,
            resolved_at=created + timedelta(hours=hours)
            if status == "resolved"
            else None,
        )

    mine = ticket(1, boss)
    second = ticket(2, boss, created_days=40, hours=30)
    open_one = ticket(3, other_user, status="in_progress", category=None)
    session.add_all([mine, second, open_one])
    await session.flush()
    session.add_all(
        [
            TicketAssignment(
                request_id=mine.id,
                technician_id=technician.id,
                assigned_by_user_id=boss.id,
            ),
            TicketAssignment(
                request_id=open_one.id,
                technician_id=technician.id,
                assigned_by_user_id=boss.id,
            ),
            MaintenanceRecord(
                equipment_id=equipment.id,
                request_id=mine.id,
                technician_id=technician.id,
                work_type="corrective",
                started_at=NOW,
                diagnosis="d",
                actions_taken="a",
                result="ok",
                labor_cost=Decimal("100.00"),
                parts_cost=Decimal("25.50"),
            ),
            MaintenanceRecord(
                equipment_id=equipment.id,
                request_id=mine.id,
                technician_id=technician.id,
                work_type="corrective",
                started_at=NOW,
                diagnosis="d",
                actions_taken="a",
                result="ok",
                labor_cost=Decimal("10.00"),
                parts_cost=None,
            ),
        ]
    )
    await session.flush()
    return SimpleNamespace(
        boss=boss, tech_user=tech_user, other=other_user, technician=technician
    )


async def test_ticket_rows_carry_every_field_from_the_database(db_session):
    await seed_tickets(db_session)
    repo = ExportRepository(db_session)
    rows = await repo.export_tickets(ExportFilters())
    assert [r["ticket_number"] for r in rows] == [
        "T-3",
        "T-1",
        "T-2",
    ]  # newest first, ties by id
    first = rows[1]  # T-1
    assert (first["equipment"], first["serial"], first["equipment_type"]) == (
        "LT-1",
        "SN-1",
        "Laptop",
    )
    assert (first["priority"], first["sla_hours"], first["category"]) == (
        "Alta",
        24,
        "Hardware",
    )
    assert (first["requested_by"], first["technician"]) == ("Ana", "Luis")
    assert first["labor_cost"] == Decimal("110.00") and first["parts_cost"] == Decimal(
        "25.50"
    )
    assert rows[0]["category"] is None and rows[2]["technician"] is None
    assert rows[0]["labor_cost"] == 0
    assert await repo.count_tickets(ExportFilters()) == 3


async def test_filters_and_scope_apply_in_sql(db_session):
    seeded = await seed_tickets(db_session)
    repo = ExportRepository(db_session)

    async def numbers(filters=None, **scope):
        rows = await repo.export_tickets(filters or ExportFilters(), **scope)
        count = await repo.count_tickets(filters or ExportFilters(), **scope)
        assert count == len(rows)
        return sorted(r["ticket_number"] for r in rows)

    assert await numbers(ExportFilters(statuses=["resolved"])) == ["T-1", "T-2"]
    assert await numbers(
        ExportFilters(statuses=["in_progress"], priorities=["high"])
    ) == ["T-3"]
    assert await numbers(ExportFilters(categories=["hardware"])) == ["T-1", "T-2"]
    assert await numbers(ExportFilters(technician_ids=[seeded.technician.id])) == [
        "T-1",
        "T-3",
    ]
    assert await numbers(ExportFilters(technician_ids=[999])) == []
    assert await numbers(ExportFilters(date_from=NOW - timedelta(days=10))) == [
        "T-1",
        "T-3",
    ]
    assert await numbers(ExportFilters(date_to=NOW - timedelta(days=10))) == ["T-2"]
    # roles: a technician sees what is assigned to them, anyone else only their own
    assert await numbers(technician_id=seeded.technician.id) == ["T-1", "T-3"]
    assert await numbers(requester_id=seeded.boss.id) == ["T-1", "T-2"]
    assert await numbers(requester_id=seeded.other.id) == ["T-3"]
    assert await repo.technician_id_of(seeded.tech_user.id) == seeded.technician.id
    assert await repo.technician_id_of(seeded.boss.id) is None


async def test_a_real_export_end_to_end(db_session):
    await seed_tickets(db_session)
    service = ExportService(ExportRepository(db_session), FakeAudit())
    result = await service.export(
        ADMIN,
        "tickets",
        "xlsx",
        ExportRequest(filters=ExportFilters(statuses=["resolved"])),
        now=NOW,
    )
    sheet = load_workbook(io.BytesIO(result.content))["Detalle"]
    headers = [c.value for c in sheet[8]]
    assert sheet.cell(
        row=9, column=headers.index("Costo total") + 1
    ).value == pytest.approx(135.5)
    assert sheet.cell(row=10, column=1).value == "T-2"
    pdf = await service.export(ADMIN, "tickets", "pdf", ExportRequest(), now=NOW)
    assert pdf_pages(pdf.content) >= 1
