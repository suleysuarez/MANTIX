import io
from datetime import UTC, date, datetime, timedelta
from decimal import Decimal
from types import SimpleNamespace

import pytest
from fastapi import UploadFile
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app.api import providers
from app.api.dependencies import get_authenticated_user
from app.core.authz import AuthContext
from app.core.exceptions import (
    PayloadTooLargeError,
    UnprocessableError,
    UnsupportedMediaTypeError,
)
from app.core.tabular import MAX_IMPORT_ROWS, parse_table
from app.core.uploads import read_limited, safe_filename, validate_content
from app.main import app
from app.schemas.report import ImportResult
from app.services.data_import import EquipmentImportService
from app.services.exporters import (
    Column,
    ReportTable,
    get_exporter,
    neutralise_formula,
    supported_formats,
)
from app.services.report import ExportedFile, ReportService
from tests.support import perm_objects

COLUMNS = (Column("name", "Name"), Column("when", "When"), Column("cost", "Cost"))
WHEN = datetime(2026, 10, 11, 15, 30, tzinfo=UTC)
ROWS = [
    {"name": "=cmd|' /C calc'!A0", "when": WHEN, "cost": Decimal("12.5")},
    {"name": 'Café, "quoted"', "when": None, "cost": None},
    {"name": "<b>bold & co</b>", "when": date(2026, 1, 2), "cost": 3},
]
TABLE = ReportTable("Test report", COLUMNS, ROWS)

# -- exporters ----------------------------------------------------------------


def test_supported_formats_and_unknown_format():
    assert supported_formats() == ["csv", "pdf", "xlsx"]
    with pytest.raises(UnsupportedMediaTypeError):
        get_exporter("docx")


@pytest.mark.parametrize(
    ("value", "expected"),
    [
        ("=1+1", "'=1+1"),
        ("+x", "'+x"),
        ("-x", "'-x"),
        ("@x", "'@x"),
        ("plain", "plain"),
    ],
)
def test_formula_prefixes_are_neutralised(value, expected):
    assert neutralise_formula(value) == expected


def test_csv_has_bom_escapes_formulas_and_quotes_values():
    data = get_exporter("csv").render(TABLE)
    assert data.startswith(b"\xef\xbb\xbf")
    text = data.decode("utf-8-sig")
    lines = text.split("\r\n")
    assert lines[0] == "Name,When,Cost"
    assert lines[1] == "'=cmd|' /C calc'!A0,2026-10-11 15:30,12.50"
    assert '"Café, ""quoted"""' in text
    assert "2026-01-02" in text


def test_xlsx_is_a_valid_workbook_with_safe_cells():
    data = get_exporter("xlsx").render(TABLE)
    sheet = load_workbook(io.BytesIO(data)).active
    assert [c.value for c in sheet[1]] == ["Name", "When", "Cost"]
    assert sheet["A2"].value == "'=cmd|' /C calc'!A0"  # stored as text, not a formula
    assert sheet["A2"].data_type == "s"
    assert sheet["C2"].value == 12.5
    assert sheet["B3"].value is None
    assert sheet.freeze_panes == "A2"


def test_pdf_renders_a_document_even_with_markup_in_values():
    data = get_exporter("pdf").render(TABLE)
    assert data.startswith(b"%PDF-")
    assert len(data) > 1000


def test_pdf_handles_an_empty_table():
    empty = ReportTable("Empty", COLUMNS, [])
    assert get_exporter("pdf").render(empty).startswith(b"%PDF-")
    assert get_exporter("csv").render(empty).decode("utf-8-sig") == "Name,When,Cost\r\n"


# -- uploads ------------------------------------------------------------------


@pytest.mark.parametrize(
    ("raw", "expected"),
    [
        ("../../etc/passwd", "passwd"),
        ("C:\\Users\\bob\\report.xlsx", "report.xlsx"),
        ("na<me>?.csv", "name.csv"),
        ("  ..hidden. ", "hidden"),
        (None, "file"),
        ("", "file"),
        ("a" * 300 + ".csv", ("a" * 200)),
    ],
)
def test_safe_filename(raw, expected):
    assert safe_filename(raw) == expected


def make_upload(content: bytes, name="data.csv") -> UploadFile:
    return UploadFile(file=io.BytesIO(content), filename=name)


async def test_read_limited_accepts_small_and_rejects_large_files():
    assert await read_limited(make_upload(b"abc"), max_bytes=10) == b"abc"
    with pytest.raises(PayloadTooLargeError):
        await read_limited(make_upload(b"x" * 11), max_bytes=10)


ALLOWED = frozenset({".csv", ".xlsx", ".pdf"})


@pytest.mark.parametrize(
    ("name", "content"),
    [
        ("evil.exe", b"MZ..."),  # extension not allowed
        ("data.csv", b""),  # empty
        ("data.csv", b"a,b\x00c"),  # binary in a text file
        ("book.xlsx", b"not a zip"),  # magic mismatch
        ("doc.pdf", b"<html>"),  # magic mismatch
        ("noext", b"a,b"),
    ],
)
def test_validate_content_rejects_bad_files(name, content):
    with pytest.raises(UnsupportedMediaTypeError):
        validate_content(name, content, ALLOWED)


def test_validate_content_returns_the_canonical_media_type():
    assert validate_content("a.csv", b"a,b\n1,2", ALLOWED) == "text/csv"
    assert validate_content("A.PDF", b"%PDF-1.7", ALLOWED) == "application/pdf"


# -- tabular parsing ------------------------------------------------------------


def xlsx_bytes(rows):
    book = Workbook()
    for row in rows:
        book.active.append(row)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def test_parse_csv_with_semicolons_bom_and_header_normalisation():
    content = (
        "\ufeffAsset Tag;Equipment-Type;Brand\r\nLT-1;LAPTOP;Dell\r\n\r\n".encode()
    )
    assert parse_table(".csv", content) == [
        {"asset_tag": "LT-1", "equipment_type": "LAPTOP", "brand": "Dell"}
    ]


def test_parse_xlsx_converts_dates_and_numbers_to_text():
    content = xlsx_bytes(
        [
            ["asset_tag", "purchase_date", "purchase_cost"],
            ["LT-1", date(2026, 1, 2), 99.5],
        ]
    )
    assert parse_table(".xlsx", content) == [
        {"asset_tag": "LT-1", "purchase_date": "2026-01-02", "purchase_cost": "99.5"}
    ]


@pytest.mark.parametrize(
    ("content", "error"),
    [
        (b"only,a,header\n", UnprocessableError),
        (b"a,a\n1,2\n", UnprocessableError),  # duplicate columns
        (b"\xff\xfe\x00bad", UnsupportedMediaTypeError),  # not UTF-8
    ],
)
def test_parse_csv_rejects_malformed_files(content, error):
    with pytest.raises(error):
        parse_table(".csv", content)


def test_parse_rejects_too_many_rows_and_bad_workbooks():
    big = "a\n" + "x\n" * (MAX_IMPORT_ROWS + 1)
    with pytest.raises(UnprocessableError, match="too many rows"):
        parse_table(".csv", big.encode())
    with pytest.raises(UnsupportedMediaTypeError):
        parse_table(".xlsx", b"PK\x03\x04garbage")
    with pytest.raises(UnsupportedMediaTypeError):
        parse_table(".pdf", b"x")


# -- import service (fake repository) -----------------------------------------------


class FakeEquipmentRepository:
    def __init__(self, taken_tags=(), taken_serials=()):
        self.taken_tags = set(taken_tags)
        self.taken_serials = set(taken_serials)
        self.created = []
        self.status_changes = []
        self.committed = False

    async def active_type_ids_by_code(self):
        return {"LAPTOP": 1, "MONITOR": 2}

    async def find_existing_identifiers(self, tags, serials):
        return tags & self.taken_tags, serials & self.taken_serials

    async def create(self, payload):
        item = SimpleNamespace(id=len(self.created) + 1, **payload.model_dump())
        self.created.append(item)
        return item

    def record_status_change(self, **kwargs):
        self.status_changes.append(kwargs)

    async def commit(self):
        self.committed = True


class FakeAudit:
    def __init__(self):
        self.records = []

    def record(self, **kwargs):
        self.records.append(kwargs)


ACTOR = AuthContext.for_roles(SimpleNamespace(id=3), "ADMIN")


def good_row(tag, **extra):
    return {
        "asset_tag": tag,
        "equipment_type": "laptop",
        "brand": "Dell",
        "model": "X1",
        **extra,
    }


def import_service(**kwargs):
    repo = FakeEquipmentRepository(**kwargs)
    return EquipmentImportService(repo, FakeAudit()), repo


async def test_import_creates_all_rows_and_audits_once():
    service, repo = import_service()
    rows = [
        good_row(
            "LT-1", serial_number="S1", purchase_cost="10.5", purchase_date="2026-01-02"
        ),
        good_row("LT-2", status="in_repair"),
    ]
    result = await service.import_equipment(rows, ACTOR, dry_run=False)
    assert (result.success, result.created, result.total_rows) == (True, 2, 2)
    assert [c.asset_tag for c in repo.created] == ["LT-1", "LT-2"]
    assert repo.created[0].equipment_type_id == 1  # code matched case-insensitively
    assert repo.created[0].purchase_cost == Decimal("10.5")
    assert len(repo.status_changes) == 2 and repo.committed
    assert service.audit.records[0]["new_values"] == {"created": 2}


async def test_dry_run_validates_without_writing():
    service, repo = import_service()
    result = await service.import_equipment([good_row("LT-1")], ACTOR, dry_run=True)
    assert result.success and result.dry_run and result.created == 0
    assert not repo.created and not repo.committed


async def test_any_invalid_row_aborts_the_whole_import_with_row_numbers():
    service, repo = import_service(taken_tags={"LT-9"}, taken_serials={"SER-9"})
    rows = [
        good_row("LT-1"),  # row 2: fine
        good_row("LT-2", equipment_type="TOASTER"),  # row 3
        good_row("LT-9"),  # row 4: exists in DB
        good_row("LT-1"),  # row 5: repeats row 2
        good_row("LT-3", serial_number="SER-9"),  # row 6: serial exists
        good_row("LT-4", purchase_cost="-5", purchase_date="not-a-date"),  # row 7
        good_row(
            "LT-5", warranty_start_date="2026-05-01", warranty_end_date="2026-01-01"
        ),
    ]
    result = await service.import_equipment(rows, ACTOR, dry_run=False)
    assert not result.success and result.created == 0
    assert not repo.created and not repo.committed
    found = {(e.row, e.field) for e in result.errors}
    assert {
        (3, "equipment_type"),
        (4, "asset_tag"),
        (5, "asset_tag"),
        (6, "serial_number"),
        (7, "purchase_cost"),
        (7, "purchase_date"),
    } <= found
    assert any(e.row == 8 for e in result.errors)  # warranty dates
    assert all(e.row != 2 for e in result.errors)


async def test_missing_required_columns_are_reported_once():
    service, _ = import_service()
    result = await service.import_equipment(
        [{"asset_tag": "LT-1"}], ACTOR, dry_run=False
    )
    assert not result.success
    assert len(result.errors) == 1
    assert "equipment_type" in result.errors[0].message


# -- HTTP ---------------------------------------------------------------------------


class StubImportService:
    def __init__(self):
        self.calls = []

    async def import_equipment(self, rows, actor, *, dry_run):
        self.calls.append((rows, dry_run))
        return ImportResult(
            dry_run=dry_run, success=True, total_rows=len(rows), created=0, errors=[]
        )


class StubReportService:
    def __init__(self):
        self.calls = []

    async def equipment_report(self, **kwargs):
        self.calls.append(("equipment", kwargs))
        return ExportedFile(
            b"a,b", "equipment-20261011-1530.csv", "text/csv; charset=utf-8"
        )

    async def ticket_report(self, **kwargs):
        self.calls.append(("tickets", kwargs))
        return ExportedFile(b"%PDF-x", "tickets.pdf", "application/pdf")


def user_with(*permissions):
    return SimpleNamespace(
        id=1,
        roles=[
            SimpleNamespace(
                code="ADMIN",
                is_active=True,
                permissions=perm_objects(*permissions),
            )
        ],
    )


@pytest.fixture
def client():
    imports, reports = StubImportService(), StubReportService()
    app.dependency_overrides[get_authenticated_user] = lambda: user_with(
        "equipment.manage", "reports.read"
    )
    app.dependency_overrides[providers.get_equipment_import_service] = lambda: imports
    app.dependency_overrides[providers.get_report_service] = lambda: reports
    with TestClient(app, raise_server_exceptions=False) as test_client:
        test_client.imports, test_client.reports = imports, reports
        yield test_client
    app.dependency_overrides.clear()


def test_import_endpoint_parses_a_csv_upload(client):
    csv_bytes = b"asset_tag,equipment_type,brand,model\nLT-1,LAPTOP,Dell,X1\n"
    response = client.post(
        "/imports/equipment?dry_run=true",
        files={"file": ("../../inventory.csv", csv_bytes, "text/csv")},
    )
    assert response.status_code == 200
    assert response.json()["total_rows"] == 1
    rows, dry_run = client.imports.calls[0]
    assert dry_run is True and rows[0]["asset_tag"] == "LT-1"


def test_import_endpoint_rejects_wrong_type_oversize_and_spoofed_content(client):
    wrong_ext = client.post(
        "/imports/equipment", files={"file": ("a.exe", b"MZ", "text/csv")}
    )
    assert wrong_ext.status_code == 415
    assert wrong_ext.json()["code"] == "unsupported_media_type"
    spoofed = client.post(
        "/imports/equipment",
        files={"file": ("a.xlsx", b"just text", "application/octet-stream")},
    )
    assert spoofed.status_code == 415
    too_big = client.post(
        "/imports/equipment",
        files={"file": ("a.csv", b"a,b\n" + b"1,2\n" * (2 * 1024 * 1024), "text/csv")},
    )
    assert too_big.status_code == 413
    assert too_big.json()["code"] == "payload_too_large"
    assert client.imports.calls == []


def test_import_endpoint_reports_unparseable_content_as_422(client):
    response = client.post(
        "/imports/equipment", files={"file": ("a.csv", b"only_a_header\n", "text/csv")}
    )
    assert response.status_code == 422
    assert response.json()["code"] == "unprocessable"


def test_report_endpoints_send_downloads_with_safe_headers(client):
    response = client.get("/reports/equipment?format=csv&status=assigned")
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("text/csv")
    assert (
        response.headers["content-disposition"]
        == 'attachment; filename="equipment-20261011-1530.csv"'
    )
    assert response.headers["cache-control"] == "no-store"
    assert client.reports.calls[0] == (
        "equipment",
        {"export_format": "csv", "status": "assigned", "equipment_type_id": None},
    )
    pdf = client.get("/reports/tickets?format=pdf&priority=high")
    assert pdf.content.startswith(b"%PDF")


def test_report_rejects_unknown_format(client):
    assert client.get("/reports/equipment?format=docx").status_code == 422


# -- real report service over the real exporters ---------------------------------------


class FakeReportRepository:
    async def equipment_rows(self, **_):
        return [{"asset_tag": "LT-1", "brand": "Dell", "status": "assigned"}]

    async def ticket_rows(self, **_):
        return [{"ticket_number": "TCK-1", "title": "=HYPERLINK(1)"}]


class FakeDashboard:
    async def summary(self, actor, *, days):
        from app.services.dashboard import DashboardService
        from tests.test_dashboard import FakeDashboardRepository

        return await DashboardService(FakeDashboardRepository()).summary(
            actor, days=days
        )

    async def resolution_times(self, actor, *, days):
        from app.services.dashboard import DashboardService
        from tests.test_dashboard import FakeDashboardRepository

        return await DashboardService(FakeDashboardRepository()).resolution_times(
            actor, days=days
        )

    async def top_failing_equipment(self, *, days, limit):
        from app.services.dashboard import DashboardService
        from tests.test_dashboard import FakeDashboardRepository

        return await DashboardService(FakeDashboardRepository()).top_failing_equipment(
            days=days, limit=limit
        )


async def test_report_service_renders_each_format():
    service = ReportService(FakeReportRepository(), FakeDashboard())
    csv_file = await service.equipment_report(export_format="csv")
    assert csv_file.filename.startswith("equipment-") and csv_file.filename.endswith(
        ".csv"
    )
    assert b"LT-1" in csv_file.content
    tickets = await service.ticket_report(export_format="xlsx")
    assert (
        load_workbook(io.BytesIO(tickets.content)).active["B2"].value
        == "'=HYPERLINK(1)"
    )
    summary = await service.summary_report(ACTOR, export_format="pdf", days=30)
    assert (
        summary.content.startswith(b"%PDF-") and summary.media_type == "application/pdf"
    )
    with pytest.raises(UnsupportedMediaTypeError):
        await service.equipment_report(export_format="docx")


# -- report repository (real database) -----------------------------------------------------


async def test_report_repository_joins_location_custodian_and_technician(db_session):
    from app.models.auth import User
    from app.models.equipment import (
        Equipment,
        EquipmentCustodyHistory,
        EquipmentLocationHistory,
        EquipmentType,
        Location,
    )
    from app.models.maintenance import MaintenanceRequest, TicketAssignment
    from app.models.priority import Priority
    from app.models.technician import Technician
    from app.repositories.report import ReportRepository

    session = db_session
    etype = EquipmentType(code="LAP", name="Laptop", is_active=True)
    users = [
        User(
            clerk_user_id=f"c{i}",
            email=f"u{i}@x.co",
            full_name=f"User {i}",
            is_active=True,
        )
        for i in range(2)
    ]
    session.add_all(
        [etype, *users, Priority(code="high", name="High", level=3, is_active=True)]
    )
    await session.flush()
    tech = Technician(user_id=users[1].id)
    laptop = Equipment(
        asset_tag="LT-1",
        equipment_type_id=etype.id,
        brand="Dell",
        model="X",
        status="assigned",
    )
    other = Equipment(
        asset_tag="LT-2",
        equipment_type_id=etype.id,
        brand="HP",
        model="Y",
        status="in_storage",
    )
    room = Location(code="R1", name="Room 1", kind="room")
    session.add_all([tech, laptop, other, room])
    await session.flush()
    created = datetime.now(UTC) - timedelta(hours=10)
    ticket = MaintenanceRequest(
        ticket_number="T-1",
        equipment_id=laptop.id,
        requested_by_user_id=users[0].id,
        title="t",
        description="d",
        priority="high",
        status="resolved",
        created_at=created,
        resolved_at=created + timedelta(hours=4),
    )
    session.add_all(
        [
            EquipmentLocationHistory(equipment_id=laptop.id, location_id=room.id),
            EquipmentCustodyHistory(
                equipment_id=laptop.id, assigned_to_user_id=users[0].id
            ),
            ticket,
        ]
    )
    await session.flush()
    session.add(TicketAssignment(request_id=ticket.id, technician_id=tech.id))
    await session.flush()

    repo = ReportRepository(session)
    equipment = await repo.equipment_rows()
    assert [r["asset_tag"] for r in equipment] == ["LT-1", "LT-2"]
    assert (
        equipment[0]["location"] == "Room 1" and equipment[0]["custodian"] == "User 0"
    )
    assert equipment[1]["location"] is None
    only_stored = await repo.equipment_rows(status="in_storage")
    assert [r["asset_tag"] for r in only_stored] == ["LT-2"]

    tickets = await repo.ticket_rows()
    assert (
        tickets[0]["technician"] == "User 1" and tickets[0]["requested_by"] == "User 0"
    )
    assert float(tickets[0]["hours_to_resolve"]) == pytest.approx(4.0, abs=0.1)
    assert await repo.ticket_rows(priority="low") == []
    assert await repo.ticket_rows(date_from=datetime.now(UTC) + timedelta(days=1)) == []
