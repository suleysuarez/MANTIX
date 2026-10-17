"""Excel template for the inventory import, with the same identity as the report exports.

Sheet 1 («Plantilla») is the one the importer reads: a coloured header in row 1 and empty,
pre-formatted rows below, with drop-downs and number rules so most mistakes cannot be typed.
It deliberately holds no example row (a forgotten example would be imported). Sheet 2
(«Instrucciones») explains every column, shows an example and states the limits.
"""

import io
from collections.abc import Sequence

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.drawing.image import Image
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from app.core.tabular import MAX_IMPORT_ROWS
from app.exports import theme

FONT = "Calibri"
#: (column, required, width, what to write, example, rule)
COLUMNS: tuple[tuple[str, bool, int, str, str, str], ...] = (
    (
        "sku",
        True,
        16,
        "Código único del ítem",
        "SSD-512",
        "3 a 40 caracteres: letras, números, . _ -. Se guarda en MAYÚSCULAS y no puede repetirse.",
    ),
    (
        "name",
        True,
        34,
        "Nombre del ítem",
        "SSD 512 GB",
        "3 a 200 caracteres. Se guarda en formato título (Cable de Red).",
    ),
    (
        "category",
        False,
        18,
        "Categoría",
        "Repuesto",
        "Una de la lista (el nombre o su código). Si queda vacía será «Repuesto».",
    ),
    (
        "unit",
        False,
        12,
        "Unidad de medida",
        "unidad",
        "Hasta 20 caracteres. Se guarda en minúsculas. Por defecto «unidad».",
    ),
    (
        "unit_cost",
        False,
        13,
        "Costo unitario",
        "64.90",
        "Número igual o mayor a 0, con máximo 2 decimales.",
    ),
    (
        "stock",
        False,
        10,
        "Stock inicial",
        "12",
        "Entero igual o mayor a 0. Queda registrado como «Carga inicial».",
    ),
    (
        "min_stock",
        False,
        12,
        "Stock mínimo",
        "5",
        "Entero igual o mayor a 0. No puede ser mayor que el máximo.",
    ),
    (
        "max_stock",
        False,
        12,
        "Stock máximo",
        "40",
        "Entero igual o mayor al mínimo. Opcional.",
    ),
    (
        "location",
        False,
        20,
        "Ubicación",
        "Bodega A",
        "Hasta 100 caracteres. Se guarda en formato título.",
    ),
    (
        "supplier",
        False,
        24,
        "Proveedor",
        "Proveedor Andino",
        "Hasta 150 caracteres. Se guarda en formato título.",
    ),
)
ROWS_FORMATTED = 500


def _fill(color: str) -> PatternFill:
    return PatternFill("solid", fgColor=color)


def build_import_template(categories: Sequence[str]) -> bytes:
    """The workbook as bytes; `categories` are the names offered in the category drop-down."""
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "Plantilla"
    thin = Side(style="thin", color=theme.GREY)

    for index, (name, required, width, hint, _example, rule) in enumerate(
        COLUMNS, start=1
    ):
        cell = sheet.cell(row=1, column=index, value=name)
        cell.font = Font(
            name=FONT, size=11, bold=True, color=theme.WHITE if required else theme.INK
        )
        cell.fill = _fill(theme.DEEP if required else theme.ACCENT_SKY)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(bottom=Side(style="medium", color=theme.MUSTARD))
        cell.comment = Comment(
            f"{hint}{' (obligatoria)' if required else ''}\n{rule}",
            "Mantix",
            width=300,
            height=110,
        )
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.row_dimensions[1].height = 26
    sheet.freeze_panes = "A2"

    # Number formats and light zebra on the rows people will fill in.
    formats = {"unit_cost": "0.00", "stock": "0", "min_stock": "0", "max_stock": "0"}
    for row in range(2, ROWS_FORMATTED + 2):
        for index, (name, *_rest) in enumerate(COLUMNS, start=1):
            cell = sheet.cell(row=row, column=index)
            cell.font = Font(name=FONT, size=10, color=theme.INK)
            cell.border = Border(bottom=thin)
            if row % 2 == 0:
                cell.fill = _fill(theme.ZEBRA)
            if name in formats:
                cell.number_format = formats[name]
                cell.alignment = Alignment(horizontal="right")

    last = ROWS_FORMATTED + 1

    def add_rule(column: str, validation: DataValidation, message: str) -> None:
        validation.errorTitle = "Valor no válido"
        validation.error = message
        validation.showErrorMessage = True
        validation.add(f"{column}2:{column}{last}")
        sheet.add_data_validation(validation)

    letters = {
        name: get_column_letter(i) for i, (name, *_r) in enumerate(COLUMNS, start=1)
    }
    add_rule(
        letters["unit_cost"],
        DataValidation(
            type="decimal", operator="between", formula1="0", formula2="1000000"
        ),
        "Escribe un número entre 0 y 1.000.000.",
    )
    for name in ("stock", "min_stock", "max_stock"):
        add_rule(
            letters[name],
            DataValidation(
                type="whole", operator="between", formula1="0", formula2="1000000"
            ),
            "Escribe un entero entre 0 y 1.000.000.",
        )
    add_rule(
        letters["sku"],
        DataValidation(
            type="textLength", operator="between", formula1="3", formula2="40"
        ),
        "El código debe tener de 3 a 40 caracteres.",
    )
    add_rule(
        letters["name"],
        DataValidation(
            type="textLength", operator="between", formula1="3", formula2="200"
        ),
        "El nombre debe tener de 3 a 200 caracteres.",
    )

    # The category list lives on a hidden sheet (a drop-down formula is limited to 255 characters).
    lists = book.create_sheet("Listas")
    for index, name in enumerate(categories, start=1):
        lists.cell(row=index, column=1, value=name)
    lists.sheet_state = "hidden"
    if categories:
        category = DataValidation(
            type="list", formula1=f"=Listas!$A$1:$A${len(categories)}", allow_blank=True
        )
        category.errorStyle = "warning"
        add_rule(letters["category"], category, "Elige una categoría de la lista.")

    _instructions(book)
    book.move_sheet("Listas", offset=len(book.sheetnames))
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def _instructions(book: Workbook) -> None:
    sheet = book.create_sheet("Instrucciones", 1)
    sheet.sheet_view.showGridLines = False
    image = Image(str(theme.LOGO_PATH))
    image.height = 44
    image.width = int(44 * theme.LOGO_SIZE[0] / theme.LOGO_SIZE[1])
    sheet.add_image(image, "A1")
    sheet.merge_cells("C1:F2")
    title = sheet.cell(row=1, column=3, value="PLANTILLA DE IMPORTACIÓN · INVENTARIO")
    title.font = Font(name=FONT, size=16, bold=True, color=theme.DEEP)
    title.alignment = Alignment(horizontal="right", vertical="center")
    for column in range(1, 7):
        sheet.cell(row=3, column=column).border = Border(
            bottom=Side(style="medium", color=theme.ACCENT_SKY)
        )

    notes = (
        f"1. Llena la hoja «Plantilla» a partir de la fila 2 (la fila 1 son los títulos: no la cambies). Hasta {MAX_IMPORT_ROWS:,} filas.".replace(
            ",", "."
        ),
        "2. Solo las columnas sku y name son obligatorias. Puedes borrar las columnas opcionales que no uses.",
        "3. Guarda el archivo como .xlsx (o expórtalo a .csv) y súbelo en Inventario → Importar. Tamaño máximo: 5 MB.",
        "4. Antes de importar verás un resumen con las filas válidas y las que tienen errores, cada una con su fila y columna.",
        "5. Los textos se normalizan al guardar (espacios sobrantes, mayúsculas y minúsculas); «Bodega A» y «bodega a» son el mismo valor.",
    )
    for offset, text in enumerate(notes):
        cell = sheet.cell(row=5 + offset, column=1, value=text)
        cell.font = Font(name=FONT, size=10, color=theme.INK)
        sheet.merge_cells(
            start_row=5 + offset, start_column=1, end_row=5 + offset, end_column=6
        )
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        sheet.row_dimensions[5 + offset].height = 28

    header_row = 11
    for index, text in enumerate(
        ("Columna", "Obligatoria", "Qué escribir", "Ejemplo", "Regla"), start=1
    ):
        cell = sheet.cell(row=header_row, column=index, value=text)
        cell.font = Font(name=FONT, size=10, bold=True, color=theme.WHITE)
        cell.fill = _fill(theme.DEEP)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for offset, (name, required, _w, hint, example, rule) in enumerate(
        COLUMNS, start=1
    ):
        values = (name, "Sí" if required else "No", hint, example, rule)
        for index, value in enumerate(values, start=1):
            cell = sheet.cell(row=header_row + offset, column=index, value=value)
            cell.font = Font(
                name=FONT, size=10, bold=required and index == 1, color=theme.INK
            )
            cell.alignment = Alignment(
                horizontal="center" if index in (1, 2, 4) else "left",
                vertical="center",
                wrap_text=True,
            )
            if offset % 2 == 0:
                cell.fill = _fill(theme.ZEBRA)
            if index == 2 and required:
                cell.fill = _fill(theme.MUSTARD)
        sheet.row_dimensions[header_row + offset].height = 32
    for letter, size in zip("ABCDE", (14, 12, 26, 18, 70), strict=True):
        sheet.column_dimensions[letter].width = size
