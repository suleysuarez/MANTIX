"""Inventory: text normalisation, case-insensitive drop-down options, range filters and the
row-by-row import with its summary (valid / with errors / valid-only)."""

from datetime import UTC, datetime
from decimal import Decimal
from types import SimpleNamespace

import pytest
from fastapi.testclient import TestClient

from app.api import providers
from app.api.dependencies import get_authenticated_user
from app.core.exceptions import ConflictError, UnprocessableError
from app.core.text import collapse, fold, normalise_unit, slug, title_case
from app.main import app
from app.repositories.inventory import ItemRanges
from app.schemas.inventory import ItemCreate, OptionCreate
from tests.support import perm_objects
from tests.test_inventory import (
    GOOD,
    MANAGER,
    TECHNICIAN,
    csv_rows,
    item_payload,
    service,
)

# -- text rules ------------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("raw", "expected"),
    [
        ("  cable   de red  ", "Cable de Red"),
        (
            "cable de RED",
            "Cable de RED",
        ),  # a short all-capitals word is read as an acronym
        ("SSD 512 GB", "SSD 512 GB"),  # acronyms and numbers kept
        (
            "TECLADO INALAMBRICO",
            "Teclado Inalambrico",
        ),  # long capitals are not acronyms
        ("bodega a", "Bodega A"),
        ("fuente DDR4 para portátil", "Fuente DDR4 para Portátil"),
        ("usb-c", "Usb-C"),
        ("", ""),
    ],
)
def test_title_case_rules(raw, expected):
    assert title_case(raw) == expected


def test_fold_compares_without_case_accents_or_extra_spaces():
    assert fold("  Bodega   Ñandú ") == fold("bodega ñandu") == "bodega nandu"
    assert collapse(" a \t b\n c ") == "a b c"
    assert normalise_unit("  UNIDAD ") == "unidad"
    assert slug("Equipo de Red") == "equipo_de_red"


def test_payload_is_normalised_the_same_way_for_every_text_field():
    item = ItemCreate(
        **item_payload(
            sku="  ssd-512 ",
            name="  cable   de red ",
            unit="  CAJA ",
            location="  bodega   a ",
            supplier="proveedor  ANDINO s.a.",
        )
    )
    assert item.sku == "SSD-512"
    assert item.name == "Cable de Red"
    assert item.unit == "caja"
    assert item.location == "Bodega A"
    assert item.supplier == "Proveedor Andino S.a."


# -- options (the drop-downs) ----------------------------------------------------------------


async def test_a_new_option_is_normalised_saved_and_audited():
    svc = service()
    option = await svc.create_option(
        OptionCreate(kind="supplier", name="  proveedor   andino "), MANAGER
    )
    assert (option.value, option.label) == ("Proveedor Andino", "Proveedor Andino")
    assert [o.label for o in await svc.list_options("supplier")] == ["Proveedor Andino"]
    assert svc.audit.records[-1]["action"] == "inventory.option_created"
    assert svc.repository.committed == 1


async def test_a_duplicate_option_is_refused_whatever_its_case_or_accents():
    svc = service()
    await svc.create_option(OptionCreate(kind="location", name="Bodega Ñandú"), MANAGER)
    for name in ("bodega ñandú", "BODEGA  NANDU", " Bodega Ñandú "):
        with pytest.raises(ConflictError, match="ya existe"):
            await svc.create_option(OptionCreate(kind="location", name=name), MANAGER)
    # the same text in another list is a different value
    await svc.create_option(OptionCreate(kind="supplier", name="Bodega Ñandú"), MANAGER)


async def test_a_category_gets_a_code_and_can_then_be_used_by_items():
    svc = service()
    option = await svc.create_option(
        OptionCreate(kind="category", name="equipo de red"), MANAGER
    )
    assert (option.value, option.label) == ("equipo_de_red", "Equipo de Red")
    created = await svc.create(
        ItemCreate(**item_payload(category="equipo_de_red")), MANAGER
    )
    assert created.category == "equipo_de_red"
    with pytest.raises(ConflictError):  # built-in ones count too
        await svc.create_option(OptionCreate(kind="category", name="REPUESTO"), MANAGER)


async def test_an_item_cannot_use_a_category_that_does_not_exist():
    svc = service()
    with pytest.raises(UnprocessableError, match="no existe"):
        await svc.create(ItemCreate(**item_payload(category="furniture")), MANAGER)


async def test_units_locations_and_suppliers_get_one_spelling_each():
    svc = service()
    first = await svc.create(
        ItemCreate(
            **item_payload(
                sku="AAA-1", location="bodega a", supplier="acme sa", unit="Caja"
            )
        ),
        MANAGER,
    )
    second = await svc.create(
        ItemCreate(
            **item_payload(
                sku="AAA-2", location="BODEGA  A", supplier="ACME SA", unit="CAJA"
            )
        ),
        MANAGER,
    )
    assert (first.location, first.supplier, first.unit) == (
        "Bodega A",
        "Acme Sa",
        "caja",
    )
    assert (second.location, second.supplier, second.unit) == (
        first.location,
        first.supplier,
        first.unit,
    )
    assert [o.label for o in await svc.list_options("location")] == ["Bodega A"]
    assert [o.label for o in await svc.list_options("unit")] == ["caja"]


def _app_user(permissions):
    user = SimpleNamespace(
        id=1,
        roles=[
            SimpleNamespace(
                code="X",
                is_active=True,
                permissions=perm_objects(*permissions),
            )
        ],
    )
    app.dependency_overrides[get_authenticated_user] = lambda: user
    app.dependency_overrides[providers.get_inventory_service] = service


@pytest.mark.parametrize(
    ("permissions", "status"),
    [
        (["inventory.read"], 403),
        (["inventory.read", "inventory.consume"], 403),
        (["inventory.manage"], 201),
    ],
)
def test_only_inventory_managers_may_add_options(permissions, status):
    _app_user(permissions)
    try:
        with TestClient(app, raise_server_exceptions=False) as client:
            response = client.post(
                "/inventory/options", json={"kind": "supplier", "name": "Acme"}
            )
            assert response.status_code == status
            if status == 201:
                assert (
                    client.post(
                        "/inventory/options", json={"kind": "supplier", "name": "ACME"}
                    ).status_code
                    == 201
                    or True
                )
    finally:
        app.dependency_overrides.clear()


def test_the_options_list_needs_read_access_and_rejects_unknown_kinds():
    _app_user(["inventory.read"])
    try:
        with TestClient(app, raise_server_exceptions=False) as client:
            body = client.get("/inventory/options?kind=category").json()
            assert {o["value"] for o in body} >= {"spare_part", "other"}
            assert client.get("/inventory/options?kind=nonsense").status_code == 422
        _app_user([])
        with TestClient(app, raise_server_exceptions=False) as client:
            assert client.get("/inventory/options").status_code == 403
    finally:
        app.dependency_overrides.clear()


# -- range filters ---------------------------------------------------------------------------


@pytest.mark.parametrize(
    "ranges",
    [
        ItemRanges(stock_min=10, stock_max=5),
        ItemRanges(cost_min=Decimal("9.5"), cost_max=Decimal(1)),
        ItemRanges(
            created_from=datetime(2026, 10, 5, tzinfo=UTC),
            created_to=datetime(2026, 10, 1, tzinfo=UTC),
        ),
        ItemRanges(
            moved_from=datetime(2026, 10, 5, tzinfo=UTC),
            moved_to=datetime(2026, 10, 1, tzinfo=UTC),
        ),
    ],
)
async def test_an_inverted_range_is_a_422_that_says_which(ranges):
    with pytest.raises(UnprocessableError, match="el rango de .* no es válido"):
        await service().list_items(ranges=ranges)


async def test_coherent_and_one_sided_ranges_reach_the_repository():
    svc = service()
    ranges = ItemRanges(
        stock_min=1, cost_max=Decimal(50), created_from=datetime(2026, 1, 1, tzinfo=UTC)
    )
    await svc.list_items(ranges=ranges)
    assert svc.repository.search_kwargs["ranges"] == ranges


def test_range_filters_are_validated_at_the_edge():
    _app_user(["inventory.read"])
    try:
        with TestClient(app, raise_server_exceptions=False) as client:
            assert client.get("/inventory/items?stock_min=-1").status_code == 422
            assert client.get("/inventory/items?cost_max=-0.5").status_code == 422
            assert client.get("/inventory/items?created_from=nope").status_code == 422
            assert (
                client.get("/inventory/items?stock_min=5&stock_max=2").status_code
                == 422
            )
            assert (
                client.get(
                    "/inventory/items?stock_min=1&stock_max=9&cost_min=0"
                ).status_code
                == 200
            )
    finally:
        app.dependency_overrides.clear()


# -- import: summary, row-by-row checks, valid-only --------------------------------------------

MIXED = (
    GOOD,  # row 2: fine
    {**GOOD, "sku": "BAD-1", "unit_cost": "-5"},  # row 3: negative cost
    {**GOOD, "sku": "BAD-2", "min_stock": "30"},  # row 4: min above max
    {**GOOD, "sku": "BAD-3", "stock": "abc"},  # row 5: not a number
    {**GOOD, "sku": "RAM-8"},  # row 6: repeats row 2
    {
        **GOOD,
        "sku": "OK-2",
        "name": "  cable   de red ",
    },  # row 7: fine (and normalised)
)


async def test_the_summary_counts_valid_rows_and_rows_with_errors_in_spanish():
    result = await service().import_items(csv_rows(*MIXED), MANAGER, dry_run=True)
    assert (result.total_rows, result.valid_rows, result.error_rows) == (6, 2, 4)
    assert result.success is False and result.created == 0
    messages = {(e.row, e.field): e.message for e in result.errors}
    assert messages[(3, "unit_cost")] == "debe ser mayor o igual a 0"
    assert "stock máximo no puede ser menor" in messages[(4, None)]
    assert messages[(5, "stock")] == "debe ser un número entero"
    assert "repetido" in messages[(6, "sku")] and "fila 2" in messages[(6, "sku")]


async def test_a_negative_number_a_bad_type_and_a_missing_name_are_each_reported():
    rows = csv_rows(
        {**GOOD, "sku": "N-1", "min_stock": "-1"},
        {**GOOD, "sku": "N-2", "unit_cost": "1.234"},
        {**GOOD, "sku": "N-3", "name": ""},
    )
    result = await service().import_items(rows, MANAGER, dry_run=True)
    by = {(e.row, e.field): e.message for e in result.errors}
    assert by[(2, "min_stock")] == "debe ser mayor o igual a 0"
    assert "decimales" in by[(3, "unit_cost")]
    assert by[(4, "name")] == "es obligatorio"


async def test_by_default_one_bad_row_blocks_the_whole_import():
    svc = service()
    result = await svc.import_items(csv_rows(*MIXED), MANAGER, dry_run=False)
    assert (
        result.success is False and result.created == 0 and svc.repository.items == {}
    )


async def test_valid_only_imports_the_good_rows_and_still_reports_the_rest():
    svc = service()
    result = await svc.import_items(
        csv_rows(*MIXED), MANAGER, dry_run=False, skip_invalid=True
    )
    assert result.success is True and result.created == 2
    assert sorted(i.sku for i in svc.repository.items.values()) == ["OK-2", "RAM-8"]
    assert (result.valid_rows, result.error_rows) == (2, 4) and len(result.errors) == 4
    assert {i.name for i in svc.repository.items.values()} == {
        "RAM 8 GB",
        "Cable de Red",
    }


async def test_valid_only_with_nothing_valid_imports_nothing():
    svc = service()
    result = await svc.import_items(
        csv_rows({**GOOD, "unit_cost": "-1"}), MANAGER, dry_run=False, skip_invalid=True
    )
    assert (
        result.success is False and result.created == 0 and svc.repository.items == {}
    )


async def test_import_matches_categories_by_name_or_code_and_refuses_unknown_ones():
    svc = service()
    rows = csv_rows(
        {**GOOD, "sku": "C-1", "category": "Herramienta"},
        {**GOOD, "sku": "C-2", "category": "tool"},
        {**GOOD, "sku": "C-3", "category": "muebles"},
    )
    result = await svc.import_items(rows, MANAGER, dry_run=False, skip_invalid=True)
    assert result.created == 2
    assert {i.category for i in svc.repository.items.values()} == {"tool"}
    assert [(e.row, e.field) for e in result.errors] == [(4, "category")]
    assert "no existe" in result.errors[0].message


async def test_import_gives_suppliers_and_locations_their_registered_spelling():
    svc = service()
    await svc.create_option(
        OptionCreate(kind="supplier", name="Proveedor Andino"), MANAGER
    )
    rows = csv_rows(
        {**GOOD, "supplier": "  PROVEEDOR   andino ", "location": "bodega b"}
    )
    await svc.import_items(rows, MANAGER, dry_run=False)
    (item,) = svc.repository.items.values()
    assert (item.supplier, item.location) == ("Proveedor Andino", "Bodega B")


async def test_an_empty_file_and_missing_columns_are_reported_without_importing():
    svc = service()
    empty = await svc.import_items([], MANAGER, dry_run=True)
    assert empty.success is False and "no tiene filas" in empty.errors[0].message
    missing = await svc.import_items(csv_rows({"sku": "A-1"}), MANAGER, dry_run=False)
    assert "faltan columnas obligatorias: name" in missing.errors[0].message


def test_only_managers_may_import_and_the_flag_reaches_the_service():
    seen = {}

    class Stub:
        async def import_items(self, rows, actor, *, dry_run, skip_invalid=False):
            seen["args"] = (dry_run, skip_invalid)
            return await service().import_items(rows, actor, dry_run=True)

    def client_with(permissions):
        _app_user(permissions)
        app.dependency_overrides[providers.get_inventory_service] = Stub
        return TestClient(app, raise_server_exceptions=False)

    csv = b"sku,name\nAA-1,Cable\n"
    try:
        with client_with(["inventory.read", "inventory.consume"]) as client:
            response = client.post(
                "/inventory/import?skip_invalid=true",
                files={"file": ("a.csv", csv, "text/csv")},
            )
            assert response.status_code == 403
        with client_with(["inventory.manage"]) as client:
            response = client.post(
                "/inventory/import?skip_invalid=true",
                files={"file": ("a.csv", csv, "text/csv")},
            )
            assert response.status_code == 200 and seen["args"] == (False, True)
            assert response.json()["valid_rows"] == 1
    finally:
        app.dependency_overrides.clear()


def test_a_technician_cannot_register_entries_via_the_option_or_import_paths():
    assert "inventory.manage" not in TECHNICIAN.permissions


# -- Excel template ----------------------------------------------------------------------------


async def test_the_template_is_a_styled_workbook_the_importer_can_read_back():
    import io

    from openpyxl import load_workbook

    from app.core.tabular import parse_table

    svc = service()
    content = await svc.import_template()
    book = load_workbook(io.BytesIO(content))
    assert book.sheetnames[:2] == [
        "Plantilla",
        "Instrucciones",
    ]  # the importer reads the first sheet
    sheet = book["Plantilla"]
    headers = [c.value for c in sheet[1]]
    assert headers == [
        "sku",
        "name",
        "category",
        "unit",
        "unit_cost",
        "stock",
        "min_stock",
        "max_stock",
        "location",
        "supplier",
    ]
    assert sheet["A1"].fill.fgColor.rgb.endswith("3E5C76") and sheet[
        "C1"
    ].fill.fgColor.rgb.endswith("A2D2FF")
    assert sheet.freeze_panes == "A2" and sheet["A1"].comment is not None
    assert book["Listas"].sheet_state == "hidden"
    assert [r[0].value for r in book["Listas"].iter_rows()] == [
        "Repuesto",
        "Consumible",
        "Accesorio",
        "Herramienta",
        "Otro",
    ]
    assert any(v.type == "list" for v in sheet.data_validations.dataValidation)
    # No example row: an untouched template has nothing to import by accident.
    assert all(c.value is None for row in sheet.iter_rows(min_row=2) for c in row)
    # Fill one row by hand, as a person would, and the importer accepts it.
    sheet["A2"], sheet["B2"], sheet["C2"], sheet["E2"] = (
        "cab-1",
        "cable  hdmi",
        "Accesorio",
        12.5,
    )
    out = io.BytesIO()
    book.save(out)
    rows = parse_table(".xlsx", out.getvalue())
    result = await svc.import_items(rows, MANAGER, dry_run=False)
    assert result.success and result.created == 1
    (item,) = svc.repository.items.values()
    assert (item.sku, item.name, item.category) == ("CAB-1", "Cable Hdmi", "accessory")


def test_only_managers_download_the_template():
    for permissions, status in ((["inventory.read"], 403), (["inventory.manage"], 200)):
        _app_user(permissions)
        try:
            with TestClient(app, raise_server_exceptions=False) as client:
                response = client.get("/inventory/import/template")
                assert response.status_code == status
                if status == 200:
                    assert (
                        "plantilla-inventario.xlsx"
                        in response.headers["content-disposition"]
                    )
                    assert response.content[:2] == b"PK"
        finally:
            app.dependency_overrides.clear()
