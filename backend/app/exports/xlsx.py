"""Excel renderer (openpyxl): "Resumen" with KPIs and native charts, "Detalle" with a
filtered, frozen table and a totals row of formulas, plus any extra sheet.

Numbers, dates and money are written as real values with number formats, never as text.
"""

import io
from datetime import UTC, datetime
from typing import Any

from openpyxl import Workbook
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.series import DataPoint
from openpyxl.drawing.image import Image
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.exports import theme
from app.exports.formatting import local
from app.exports.model import ExportColumn, ReportExport
from app.services.exporters import neutralise_formula

FONT = "Calibri"
_NUMBER_FORMATS = {
    "date": theme.XLSX_DATE,
    "datetime": theme.XLSX_DATETIME,
    "money": theme.XLSX_MONEY,
    "number": theme.XLSX_NUMBER,
    "int": theme.XLSX_INTEGER,
}
_THIN = Side(style="thin", color=theme.GREY)
_HEADER_FILL = PatternFill("solid", fgColor=theme.DEEP)
_ZEBRA_FILL = PatternFill("solid", fgColor=theme.ZEBRA)
_TOTAL_FILL = PatternFill("solid", fgColor=theme.ACCENT_LIGHT)
#: Rows above the table: logo and title (1-3), metadata (4-6), a blank row (7).
HEADER_ROW = 8
LOGO_HEIGHT_PX = 44


def _cell_value(value: Any, column: ExportColumn) -> Any:
    """A value Excel can store: text is neutralised against formulas, dates are naive local time."""
    if value is None or value == "":
        return None
    if column.kind == "bool":
        return "Sí" if value else "No"
    if column.kind in ("date", "datetime") and isinstance(value, datetime):
        return local(value).replace(tzinfo=None)
    if column.kind == "money":
        return float(value)
    if isinstance(value, str):
        return neutralise_formula(value)
    return value


def _logo(sheet: Worksheet) -> None:
    image = Image(str(theme.LOGO_PATH))
    ratio = theme.LOGO_SIZE[0] / theme.LOGO_SIZE[1]
    image.height = LOGO_HEIGHT_PX
    image.width = int(LOGO_HEIGHT_PX * ratio)
    sheet.add_image(image, "A1")
    for row in (1, 2):
        sheet.row_dimensions[row].height = 20


def _title_block(
    sheet: Worksheet,
    export: ReportExport,
    title: str,
    span: int,
    right_column: int | None = None,
) -> None:
    _logo(sheet)
    last = max(span, 4)
    sheet.merge_cells(start_row=1, start_column=3, end_row=2, end_column=last)
    cell = sheet.cell(row=1, column=3, value=title)
    cell.font = Font(name=FONT, size=16, bold=True, color=theme.DEEP)
    cell.alignment = Alignment(horizontal="right", vertical="center")
    for column in range(1, last + 1):
        sheet.cell(row=3, column=column).border = Border(
            bottom=Side(style="medium", color=theme.ACCENT_SKY)
        )
    rows = max(len(export.meta_left), len(export.meta_right))
    right = right_column or max(last // 2, 2) + 1
    for index in range(rows):
        for items, column in ((export.meta_left, 1), (export.meta_right, right)):
            if index < len(items):
                label = sheet.cell(
                    row=4 + index, column=column, value=items[index].label
                )
                label.font = Font(name=FONT, size=10, bold=True, color=theme.DEEP)
                value = sheet.cell(
                    row=4 + index,
                    column=column + 1,
                    value=neutralise_formula(items[index].value),
                )
                value.font = Font(name=FONT, size=10, color=theme.INK)


def _write_table(
    sheet: Worksheet,
    columns: list[ExportColumn],
    rows: list[dict[str, Any]],
    *,
    header_row: int,
    totals_label: str = "Total",
) -> int:
    """Header, rows with real formats, autofilter, freeze panes and a totals row of formulas."""
    for index, column in enumerate(columns, start=1):
        cell = sheet.cell(row=header_row, column=index, value=column.header)
        cell.font = Font(name=FONT, size=10, bold=True, color=theme.WHITE)
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(
            horizontal="right" if column.numeric else "left",
            vertical="center",
            wrap_text=True,
        )
        sheet.column_dimensions[get_column_letter(index)].width = column.width
    sheet.row_dimensions[header_row].height = 24

    first = header_row + 1
    for offset, row in enumerate(rows):
        r = first + offset
        for index, column in enumerate(columns, start=1):
            cell = sheet.cell(
                row=r, column=index, value=_cell_value(row.get(column.key), column)
            )
            cell.font = Font(name=FONT, size=10, color=theme.INK)
            cell.alignment = Alignment(
                horizontal="right" if column.numeric else "left",
                vertical="center",
                wrap_text=column.kind == "text",
            )
            if column.kind in _NUMBER_FORMATS:
                cell.number_format = _NUMBER_FORMATS[column.kind]
            cell.border = Border(bottom=_THIN)
            if offset % 2 == 1:
                cell.fill = _ZEBRA_FILL
    last = first + len(rows) - 1
    if rows:
        sheet.auto_filter.ref = f"A{header_row}:{get_column_letter(len(columns))}{last}"
    sheet.freeze_panes = f"A{first}"

    if rows and any(column.total for column in columns):
        total_row = last + 1
        for index, column in enumerate(columns, start=1):
            letter = get_column_letter(index)
            cell = sheet.cell(row=total_row, column=index)
            if column.total == "sum":
                cell.value = f"=SUBTOTAL(109,{letter}{first}:{letter}{last})"
            elif column.total == "count":
                formula = f"SUBTOTAL(103,{letter}{first}:{letter}{last})"
                cell.value = (
                    f'="{totals_label}: "&{formula}' if index == 1 else f"={formula}"
                )
            elif index == 1:
                cell.value = totals_label
            cell.font = Font(name=FONT, size=10, bold=True, color=theme.DEEP)
            cell.fill = _TOTAL_FILL
            cell.border = Border(top=Side(style="medium", color=theme.DEEP))
            cell.alignment = Alignment(horizontal="right" if column.numeric else "left")
            if column.kind in _NUMBER_FORMATS and column.total == "sum":
                cell.number_format = _NUMBER_FORMATS[column.kind]
            elif column.total == "count":
                cell.number_format = theme.XLSX_INTEGER
    return last


def _style_chart(chart: BarChart | PieChart, title: str) -> None:
    chart.title = title
    chart.height, chart.width = 7.5, 12.5
    chart.legend = None if isinstance(chart, BarChart) else chart.legend


def _summary_sheet(book: Workbook, export: ReportExport, include_charts: bool) -> None:
    sheet = book.active
    sheet.title = "Resumen"
    sheet.sheet_view.showGridLines = False
    for letter, width in zip("ABCDEFGH", (30, 16, 4, 30, 16, 4, 4, 4), strict=True):
        sheet.column_dimensions[letter].width = width
    _title_block(sheet, export, export.title, 5, right_column=4)

    row = HEADER_ROW
    placed: list[tuple[Any, int, int, int]] = []  # group, header row, last row, column
    for pair_start in range(0, len(export.summary), 2):
        pair = export.summary[pair_start : pair_start + 2]
        for offset, group in enumerate(pair):
            column = 1 if offset == 0 else 4
            head = sheet.cell(row=row, column=column, value=group.title)
            head.font = Font(name=FONT, size=11, bold=True, color=theme.WHITE)
            head.fill = _HEADER_FILL
            sheet.cell(row=row, column=column + 1).fill = _HEADER_FILL
            for index, line in enumerate(group.lines, start=1):
                label = sheet.cell(
                    row=row + index, column=column, value=neutralise_formula(line.label)
                )
                label.font = Font(name=FONT, size=10, color=theme.INK)
                value = sheet.cell(row=row + index, column=column + 1, value=line.value)
                value.font = Font(name=FONT, size=10, bold=True, color=theme.DEEP)
                value.alignment = Alignment(horizontal="right")
                if isinstance(line.value, float):
                    value.number_format = theme.XLSX_NUMBER
                label.border = value.border = Border(bottom=_THIN)
            placed.append((group, row, row + len(group.lines), column))
        row += max(len(group.lines) for group in pair) + 3

    if include_charts:
        anchor_row = row + 1
        slot = 0
        for group, first, last, column in placed:
            if group.chart is None or last <= first:
                continue
            labels = Reference(sheet, min_col=column, min_row=first + 1, max_row=last)
            values = Reference(sheet, min_col=column + 1, min_row=first, max_row=last)
            chart: BarChart | PieChart
            if group.chart == "bar":
                chart = BarChart()
                chart.type = "col"
                chart.add_data(values, titles_from_data=False)
                chart.set_categories(labels)
                chart.series[0].graphicalProperties.solidFill = theme.DEEP
                chart.dataLabels = DataLabelList()
                chart.dataLabels.showVal = True
                chart.y_axis.majorGridlines = None
            else:
                chart = PieChart()
                chart.add_data(values, titles_from_data=False)
                chart.set_categories(labels)
                for index in range(last - first):
                    point = DataPoint(idx=index)
                    point.graphicalProperties.solidFill = theme.CHART_COLORS[
                        index % len(theme.CHART_COLORS)
                    ]
                    chart.series[0].dPt.append(point)
                chart.dataLabels = DataLabelList()
                chart.dataLabels.showPercent = True
            _style_chart(chart, group.title)
            sheet.add_chart(
                chart, f"{'A' if slot % 2 == 0 else 'D'}{anchor_row + (slot // 2) * 17}"
            )
            slot += 1


def _detail_sheet(book: Workbook, export: ReportExport) -> None:
    sheet = book.create_sheet(export.sheet_name)
    sheet.sheet_view.showGridLines = False
    columns = list(export.columns)
    _title_block(sheet, export, export.title, len(columns))
    _write_table(sheet, columns, list(export.rows), header_row=HEADER_ROW)
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_title_rows = f"{HEADER_ROW}:{HEADER_ROW}"


def _extra_sheet(book: Workbook, export: ReportExport, extra: Any) -> None:
    sheet = book.create_sheet(extra.name)
    sheet.sheet_view.showGridLines = False
    columns = list(extra.columns)
    _title_block(sheet, export, extra.title, len(columns))
    _write_table(sheet, columns, list(extra.rows), header_row=HEADER_ROW)


def render_xlsx(
    export: ReportExport,
    *,
    include_summary: bool = True,
    include_charts: bool = True,
    generated_at: datetime | None = None,
) -> bytes:
    book = Workbook()
    if include_summary and export.summary:
        _summary_sheet(book, export, include_charts)
        _detail_sheet(book, export)
    else:
        sheet = book.active
        book.remove(sheet)
        _detail_sheet(book, export)
    for extra in export.extra_sheets:
        _extra_sheet(book, export, extra)
    book.properties.title = export.title.title()
    book.properties.creator = theme.SYSTEM_NAME
    book.properties.created = (generated_at or datetime.now(UTC)).replace(tzinfo=None)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()
