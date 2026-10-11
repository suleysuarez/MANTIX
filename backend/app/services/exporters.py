"""Tabular exporters (Strategy pattern): CSV, Excel and PDF.

Each exporter turns the same `Table` into bytes, so a report is described once
and rendered in any format. Spreadsheet formats neutralise "formula injection":
text that starts with `=`, `+`, `-`, `@` is prefixed with a quote so Excel
never evaluates user-supplied content.
"""

import csv
import io
from abc import ABC, abstractmethod
from collections.abc import Sequence
from dataclasses import dataclass
from datetime import UTC, date, datetime
from decimal import Decimal
from typing import Any, ClassVar

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.core.exceptions import UnsupportedMediaTypeError

_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")
_MAX_CELL_CHARS = 32000  # Excel's hard limit is 32,767


@dataclass(frozen=True, slots=True)
class Column:
    key: str
    header: str
    width: int = 18  # approximate Excel character width


@dataclass(frozen=True, slots=True)
class ReportTable:
    """A titled grid of rows (dicts keyed by `Column.key`)."""

    title: str
    columns: Sequence[Column]
    rows: Sequence[dict[str, Any]]
    subtitle: str | None = None


def neutralise_formula(value: str) -> str:
    return "'" + value if value.startswith(_FORMULA_PREFIXES) else value


def _as_text(value: Any) -> str:
    """Plain-text rendering used by CSV and PDF."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "yes" if value else "no"
    if isinstance(value, datetime):
        return _naive_utc(value).strftime("%Y-%m-%d %H:%M")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, Decimal):
        return f"{value:.2f}"
    return str(value)


def _naive_utc(value: datetime) -> datetime:
    if value.tzinfo is not None:
        value = value.astimezone(UTC).replace(tzinfo=None)
    return value


class Exporter(ABC):
    format: ClassVar[str]
    extension: ClassVar[str]
    media_type: ClassVar[str]

    @abstractmethod
    def render(self, table: ReportTable) -> bytes: ...


class CsvExporter(Exporter):
    format = "csv"
    extension = "csv"
    media_type = "text/csv; charset=utf-8"

    def render(self, table: ReportTable) -> bytes:
        buffer = io.StringIO(newline="")
        writer = csv.writer(buffer, lineterminator="\r\n")
        writer.writerow([neutralise_formula(c.header) for c in table.columns])
        for row in table.rows:
            writer.writerow(
                [neutralise_formula(_as_text(row.get(c.key))) for c in table.columns]
            )
        # UTF-8 BOM so Excel opens accented text correctly.
        return buffer.getvalue().encode("utf-8-sig")


class XlsxExporter(Exporter):
    format = "xlsx"
    extension = "xlsx"
    media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    def render(self, table: ReportTable) -> bytes:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = table.title[:31]
        header_fill = PatternFill("solid", fgColor="1F3A5F")
        sheet.append([c.header for c in table.columns])
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
            cell.alignment = Alignment(vertical="center")
        for row in table.rows:
            sheet.append([self._cell(row.get(c.key)) for c in table.columns])
        for index, column in enumerate(table.columns, start=1):
            sheet.column_dimensions[get_column_letter(index)].width = column.width
        sheet.freeze_panes = "A2"
        if table.rows:
            sheet.auto_filter.ref = sheet.dimensions
        buffer = io.BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()

    @staticmethod
    def _cell(value: Any) -> Any:
        if value is None:
            return None
        if isinstance(value, datetime):
            return _naive_utc(value)
        if isinstance(value, bool | int | float | date):
            return value
        if isinstance(value, Decimal):
            return float(value)
        return neutralise_formula(str(value))[:_MAX_CELL_CHARS]


class PdfExporter(Exporter):
    format = "pdf"
    extension = "pdf"
    media_type = "application/pdf"

    def render(self, table: ReportTable) -> bytes:
        buffer = io.BytesIO()
        document = SimpleDocTemplate(
            buffer,
            pagesize=landscape(A4),
            leftMargin=12 * mm,
            rightMargin=12 * mm,
            topMargin=14 * mm,
            bottomMargin=14 * mm,
            title=table.title,
            author="Mantix",
        )
        styles = getSampleStyleSheet()
        cell_style = ParagraphStyle(
            "cell", parent=styles["BodyText"], fontSize=7, leading=9
        )
        head_style = ParagraphStyle(
            "head", parent=cell_style, textColor=colors.white, fontName="Helvetica-Bold"
        )
        story: list[Any] = [Paragraph(_escape(table.title), styles["Title"])]
        generated = datetime.now(UTC).strftime("%Y-%m-%d %H:%M UTC")
        subtitle = f"{table.subtitle} - " if table.subtitle else ""
        story.append(
            Paragraph(
                f"{_escape(subtitle)}Generated {generated} - {len(table.rows)} rows",
                styles["Normal"],
            )
        )
        story.append(Spacer(1, 6 * mm))

        total_width = sum(c.width for c in table.columns) or 1
        available = landscape(A4)[0] - 24 * mm
        widths = [available * c.width / total_width for c in table.columns]
        data = [[Paragraph(_escape(c.header), head_style) for c in table.columns]]
        for row in table.rows:
            data.append(
                [
                    Paragraph(_escape(_as_text(row.get(c.key))), cell_style)
                    for c in table.columns
                ]
            )
        grid = Table(data, colWidths=widths, repeatRows=1)
        grid.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F3A5F")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#BBBBBB")),
                    (
                        "ROWBACKGROUNDS",
                        (0, 1),
                        (-1, -1),
                        [colors.white, colors.HexColor("#F3F6FA")],
                    ),
                ]
            )
        )
        story.append(grid)
        document.build(story)
        return buffer.getvalue()


def _escape(text: str) -> str:
    """reportlab Paragraph parses a mini-HTML; escape user text."""
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


_EXPORTERS: dict[str, Exporter] = {
    exporter.format: exporter
    for exporter in (CsvExporter(), XlsxExporter(), PdfExporter())
}


def get_exporter(export_format: str) -> Exporter:
    try:
        return _EXPORTERS[export_format.lower()]
    except KeyError:
        raise UnsupportedMediaTypeError(
            f"unsupported export format '{export_format}'"
        ) from None


def supported_formats() -> list[str]:
    return sorted(_EXPORTERS)
