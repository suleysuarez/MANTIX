"""The Mantix document layout shared by every report: status colours, the generic builder used
by the whole gallery, and what the PDF and the workbook show."""

import io
from datetime import UTC, date, datetime
from decimal import Decimal

import pytest
from openpyxl import load_workbook
from reportlab import rl_config

from app.exports.formatting import status_tone
from app.exports.generic import build_generic_export
from app.exports.meta import meta_items, report_code
from app.exports.pdf import render_pdf
from app.exports.xlsx import render_xlsx
from app.services.report_builder import DATASETS

NOW = datetime(2026, 10, 15, 17, 45, tzinfo=UTC)  # 12:45 in Colombia
FOOTER = "MANTIX · Sistema de mantenimiento de equipos tecnológicos"


@pytest.mark.parametrize(
    ("text", "tone"),
    [
        ("Completado", "good"),
        ("Resuelto", "good"),
        ("En funcionamiento", "good"),
        ("En stock", "good"),
        ("Disponible", "good"),
        ("En proceso", "warn"),
        ("En mantenimiento", "warn"),
        ("Bajo mínimo", "warn"),
        ("Vencido", "bad"),
        ("Agotado", "bad"),
        ("Fuera de servicio", "bad"),
        ("Dado de baja", "bad"),
        ("Programado", "info"),
        ("Enviado", "info"),
        ("Pendiente", "neutral"),
        ("Inactivo", "neutral"),
        ("No disponible", "neutral"),
        ("Algo raro", "neutral"),
    ],
)
def test_a_status_is_coloured_by_what_it_says(text, tone):
    assert status_tone(text) == tone


def test_the_metadata_block_has_the_six_entries_and_no_company_data():
    left, right = meta_items(
        code=report_code("eq", NOW),
        generated_by="suley.suarez",
        generated_at=NOW,
        records=24,
        period="01/06/2026 – 16/07/2026",
        filters="Sin filtros",
    )
    labels = [m.label for m in [*left, *right]]
    assert labels == [
        "Reporte",
        "Fecha de generación",
        "Periodo",
        "Generado por",
        "Registros",
        "Filtros",
    ]
    values = {m.label: m.value for m in [*left, *right]}
    assert values["Reporte"] == "EQ-20261015-1245"
    assert values["Fecha de generación"] == "15/10/2026 12:45"
    assert not {"Empresa", "NIT", "Razón social"} & set(labels)


def _rows(key):
    """Rows shaped like the repository's, for whichever columns the dataset has."""
    samples = {
        "status": ["En funcionamiento", "En mantenimiento", "Vencido", "Completado"],
        "kind": ["Entrada", "Consumo", "Ajuste"],
        "equipment_type": ["Portátil", "Servidor"],
        "category": ["Repuesto", "Otro"],
        "priority": ["Alta", "Baja"],
        "location": ["Bodega A", "Oficina"],
        "technician": ["Ana", "Luis"],
    }
    rows = []
    for n in range(7):
        row = {}
        for column in DATASETS[key].columns:
            name = column.key
            if name in samples:
                row[name] = samples[name][n % len(samples[name])]
            elif "date" in name:
                row[name] = date(2026, 9, 1 + n)
            elif name.endswith("_at"):
                row[name] = datetime(2026, 9, 1 + n, 8, tzinfo=UTC)
            elif "cost" in name or "price" in name:
                row[name] = Decimal("10.50") * (n + 1)
            elif name in ("quantity", "stock", "units", "count", "tickets"):
                row[name] = n + 1
            else:
                row[name] = f"{name}-{n}"
        rows.append(row)
    return rows


def _document(key):
    dataset = DATASETS[key]
    return build_generic_export(
        key=key,
        title=dataset.title,
        dataset_columns=dataset.columns,
        rows=_rows(key),
        generated_by="suley.suarez",
        generated_at=NOW,
        filters="Sin filtros",
        period="Todo el historial",
    )


@pytest.mark.parametrize("key", sorted(DATASETS))
def test_every_report_of_the_gallery_gets_the_same_pdf_and_workbook(key):
    document = _document(key)
    assert document.title == f"REPORTE DE {DATASETS[key].title.upper()}"
    meta = {m.label: m.value for m in [*document.meta_left, *document.meta_right]}
    assert meta["Registros"] == "7" and meta["Generado por"] == "suley.suarez"

    assert render_pdf(document, generated_at=NOW).startswith(b"%PDF-")

    book = load_workbook(io.BytesIO(render_xlsx(document, generated_at=NOW)))
    assert book.sheetnames[0] == "Resumen" and "Detalle" in book.sheetnames
    detail = book["Detalle"]
    assert detail["C1"].value == document.title
    # clean header: logo top-left, large title beside it, rule, plain metadata, light table header
    assert detail["C1"].alignment.horizontal == "left" and detail[
        "C1"
    ].font.color.rgb.endswith("3E5C76")
    assert detail["C1"].font.size == 20 and detail["A8"].fill.fgColor.rgb.endswith(
        "ADD7F6"
    )  # light table header
    assert (
        detail["C1"].fill.patternType is None and detail["A1"].fill.patternType is None
    )
    assert detail["A3"].border.bottom.style == "medium"
    texts = {c.value for row in detail.iter_rows() for c in row if c.value}
    assert FOOTER in texts, "the sheet closes with the Mantix footer band"


def test_status_columns_become_pills_and_the_summary_counts_them():
    document = _document("equipment")
    status = next(c for c in document.columns if c.key == "status")
    assert status.pill and status.alignment == "center"
    assert not next(c for c in document.columns if c.key == "brand").pill
    cards = next(g for g in document.summary if g.kpi).lines
    assert cards[0].label == "Total de registros" and cards[0].value == 7
    assert [c.tone for c in cards[1:]] == [status_tone(c.label) for c in cards[1:]]
    assert any(g.chart == "pie" for g in document.summary)

    sheet = load_workbook(io.BytesIO(render_xlsx(document)))["Detalle"]
    column = [c.value for c in sheet[8]].index("Estado") + 1
    fills = {
        sheet.cell(row=r, column=column).value: sheet.cell(
            row=r, column=column
        ).fill.fgColor.rgb[-6:]
        for r in range(9, 16)
    }
    assert fills["Vencido"] == "FBE3E0"
    assert fills["En funcionamiento"] == "E3F5EA"
    assert fills["En mantenimiento"] == "FDF0D6"


def test_the_workbook_summary_has_cards_a_doughnut_and_a_bar_chart():
    from openpyxl.chart import BarChart, DoughnutChart

    summary = load_workbook(io.BytesIO(render_xlsx(_document("equipment"))))["Resumen"]
    kinds = [type(c) for c in summary._charts]
    assert DoughnutChart in kinds and BarChart in kinds
    assert summary.merged_cells.ranges  # the cards are merged, tinted boxes
    values = {
        c.value for row in summary.iter_rows() for c in row if c.value is not None
    }
    assert {"Total de registros", 7} <= values


def test_the_pdf_shows_metadata_summary_cards_and_footer(monkeypatch):
    monkeypatch.setattr(rl_config, "pageCompression", 0)  # readable text in the file
    pdf = render_pdf(_document("inventory_items"), generated_at=NOW).decode("latin-1")
    for expected in (
        "Generado por",
        "Registros",
        "RESUMEN GENERAL",
        "TOTAL DE REGISTROS:",
        "MANTIX \\267 Sistema de mantenimiento de equipos tecnol",
        "suley.suarez",
    ):
        assert expected in pdf, expected
    assert "NIT" not in pdf and "Raz" not in pdf
