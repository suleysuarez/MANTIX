import io
from datetime import UTC, datetime
from decimal import Decimal
from types import SimpleNamespace
from typing import Any

import pytest
from openpyxl import load_workbook
from pydantic import ValidationError
from sqlalchemy.dialects import postgresql

from app.core.exceptions import ForbiddenError
from app.exports.base import ExportContext
from app.exports.technicians import TechniciansExport, prepare
from app.repositories.export import ExportRepository
from app.schemas.export import ExportFilters, ExportRequest
from app.services.export import ExportService
from tests.test_export import ADMIN, NOW, TECHNICIAN, FakeAudit, pdf_pages


def raw_technician(n: int, **over: Any) -> dict[str, Any]:
    row = {
        "id": n,
        "full_name": f"Técnico {n}",
        "employee_code": f"EMP-{n:03d}",
        "email": f"tec{n}@mantix.co",
        "specialization": "Redes" if n % 2 else None,
        "is_active": n % 5 != 0,
        "is_available": n % 3 != 0,
        "hourly_rate": Decimal("25.50") if n % 2 else None,
        "created_at": datetime(2026, 1, n, tzinfo=UTC),
        "open_tickets": n % 4,
        "resolved": n * 2,
    }
    row.update(over)
    return row


class FakeRepository:
    def __init__(self, rows=None, technician_id=None):
        self.rows = (
            rows if rows is not None else [raw_technician(n) for n in range(1, 7)]
        )
        self._technician_id = technician_id
        self.filters: list[ExportFilters] = []
        self.runs: list[dict[str, Any]] = []

    async def organization(self):
        return None

    async def technician_id_of(self, user_id):
        return self._technician_id

    async def count_technicians(self, filters):
        self.filters.append(filters)
        return len(self.rows)

    async def export_technicians(self, filters, *, limit):
        return self.rows[:limit]

    def record_run(self, **kwargs):
        self.runs.append(kwargs)
        return SimpleNamespace(id=len(self.runs))

    async def flush(self):
        pass

    async def commit(self):
        pass


def make_service(**kwargs):
    repo = FakeRepository(**kwargs)
    return ExportService(repo, FakeAudit()), repo


def context(**filters):
    return ExportContext(
        company="Mantix",
        generated_by="Ana Pérez",
        generated_at=NOW,
        filters=ExportFilters(**filters),
    )


def test_rows_get_spanish_labels():
    active = prepare(raw_technician(1))
    assert (active["status"], active["availability"]) == ("Activo", "Disponible")
    inactive = prepare(raw_technician(5, is_available=False))
    assert (inactive["status"], inactive["availability"]) == (
        "Inactivo",
        "No disponible",
    )


async def test_excel_has_the_shared_header_with_the_title_and_no_company_ids():
    service, repo = make_service()
    out = await service.export(ADMIN, "technicians", "xlsx", ExportRequest(), now=NOW)
    assert out.filename == "mantix_technicians_2026-10-14.xlsx"
    book = load_workbook(io.BytesIO(out.content))
    sheet = book["Detalle"]
    assert sheet.cell(row=1, column=3).value == "TÉCNICOS"
    assert sheet.cell(row=3, column=1).border.bottom.style == "medium"  # divider line
    values = {c.value for row in sheet.iter_rows(min_row=4, max_row=7) for c in row}
    assert {
        "Reporte",
        "Fecha de generación",
        "Generado por",
        "Registros",
        "Filtros",
    } <= values
    assert "NIT" not in values and "Razón social" not in values
    headers = [c.value for c in sheet[8] if c.value]
    assert headers[0] == "Técnico" and "Tarifa por hora" in headers
    assert repo.runs[0]["export_format"] == "xlsx" and len(repo.runs) == 1


async def test_pdf_uses_the_compact_default_columns():
    service, _ = make_service()
    pdf = await service.export(ADMIN, "technicians", "pdf", ExportRequest(), now=NOW)
    assert pdf.media_type == "application/pdf" and pdf_pages(pdf.content) >= 1
    assert service._columns(TechniciansExport(), "pdf", None) == list(
        TechniciansExport.pdf_default
    )
    assert "email" not in TechniciansExport.pdf_default


async def test_only_admins_and_supervisors_may_export_technicians():
    for technician_id in (5, None):  # a technician and a plain requester
        service, _ = make_service(technician_id=technician_id)
        with pytest.raises(ForbiddenError):
            await service.count(TECHNICIAN, "technicians", ExportFilters())
        with pytest.raises(ForbiddenError):
            await service.export(
                TECHNICIAN, "technicians", "pdf", ExportRequest(), now=NOW
            )


def test_the_summary_and_filters_describe_the_roster():
    export = TechniciansExport().build(
        [prepare(r) for r in [raw_technician(n) for n in range(1, 4)]],
        context(is_available=True, load="low", specialty="redes"),
    )
    meta = {m.label: m.value for m in export.meta_right}
    assert "Disponibilidad: Disponible" in meta["Filtros"]
    assert (
        "Carga: 1 a 2 tickets" in meta["Filtros"]
        and "Especialidad: redes" in meta["Filtros"]
    )
    kpis = next(g for g in export.summary if g.kpi)
    assert kpis.lines[0].label == "Total de técnicos"


def test_technician_filters_are_validated_and_trimmed():
    assert ExportFilters(search="  ana  ", specialty="   ").search == "ana"
    assert ExportFilters(specialty="   ").specialty is None
    with pytest.raises(ValidationError):
        ExportFilters(load="heavy")
    with pytest.raises(ValidationError):
        ExportFilters(search="x" * 101)


@pytest.mark.parametrize("load", ["none", "low", "high"])
def test_the_roster_query_builds_for_every_filter(load):
    filters = ExportFilters(
        search="ana", specialty="red", is_active=True, is_available=False, load=load
    )
    sql = str(
        ExportRepository._technicians(None, filters).compile(  # type: ignore[arg-type]
            dialect=postgresql.dialect()
        )
    )
    assert "technicians" in sql and "open_load" in sql
    assert "users.is_active" in sql
